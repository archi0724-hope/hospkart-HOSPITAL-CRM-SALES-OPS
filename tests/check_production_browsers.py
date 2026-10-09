"""Source-independent Chrome/Edge checks through the actual Waitress WSGI server.

Install requirements-dev.txt and Chrome/Edge, then run this file from the project.
All workbooks, databases, admin hashes, and browser profiles are temporary.
"""
import json
import os
import sys
import tempfile
from pathlib import Path
from threading import Thread

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import app
from openpyxl import Workbook
from playwright.sync_api import sync_playwright
from waitress import create_server
from production import server_options
from werkzeug.security import generate_password_hash


def make_workbook(path):
    workbook=Workbook();workbook.remove(workbook.active)
    invoices=[['2026-09-01','INV-1','Test Alpha','Jaipur',100,'Paid'],['2026-09-02','INV-2','Test Alpha','Jaipur',200,'Active'],['2026-09-03','INV-3','Test Beta','Delhi',300,'Part Paid']]
    headers=['Invoice Date','Invoice No.','Client / Hospital','Delivery City / Location','Invoice Total (₹)','Invoice Status']
    summary=workbook.create_sheet('Client Summary');summary.append(headers)
    for row in invoices:summary.append(row)
    detail=workbook.create_sheet('Client Call List');detail.append(headers+['Item Name','Category','Item Total incl. GST (₹)'])
    for row,item,amount in [(invoices[0],'Item A',40),(invoices[0],'Item B',60),(invoices[1],'Item C',200),(invoices[2],'Item D',300)]:detail.append(row+[item,'Test category',amount])
    cancelled=workbook.create_sheet('Cancelled Invoices');cancelled.append(headers);cancelled.append(['2026-09-04','INV-X','Test Cancelled','Delhi',50,'Canceled'])
    for name in ['Categorized Items','Verification Summary','Read Me']:
        sheet=workbook.create_sheet(name);sheet.append(['Label','Value']);sheet.append(['Synthetic test source',name])
        if name=='Read Me':
            for number in range(8):sheet.append([f'Storage capacity fixture {number}','x'*32000])
    workbook.save(path);workbook.close()


def make_test_record():
    return {
        'id':'TEST-CLIENT-1','name':'Synthetic Test Hospital','type':'Test fixture',
        'city':'Test City','district':'Test District','state':'Test State',
        'contact':'Test Contact','mobile':'9999999999','email':'test@example.invalid',
        'requirement':'Test equipment','product':'Test product','status':'New Lead',
        'quote':0,'order':0,'last':'2026-09-01','next':'','executive':'Test',
        'remark':'Synthetic browser-test record','priority':'Medium','sourceData':{}
    }


def check():
    with tempfile.TemporaryDirectory() as temp:
        root=Path(temp);app.DATA_DIR=root;app.DB_PATH=root/'test.db';app.ADMIN_PASSWORD_PATH=root/'admin.hash'
        app.admin_attempts.clear();app.init_db()
        if app.scheduler.running:app.scheduler.pause()
        password='production-browser-test-only'
        app.ADMIN_PASSWORD_PATH.write_text(generate_password_hash(password),encoding='utf-8')
        # Ignore any real environment hash: test authorization always uses temp data.
        previous_hash=os.environ.pop('ADMIN_PASSWORD_HASH',None)
        fixture=root/'test-invoices.xlsx';make_workbook(fixture)
        options=server_options();options.update(host='127.0.0.1',port=0)
        server=create_server(app.app,**options)
        Thread(target=server.run,daemon=True).start()
        address=f'http://127.0.0.1:{server.effective_port}/'
        transfer=None;reports=[]
        try:
            with sync_playwright() as playwright:
                for channel in os.getenv('BROWSER_CHANNELS','chrome,msedge').split(','):
                    browser=playwright.chromium.launch(channel=None if channel=='chromium' else channel)
                    context=browser.new_context(viewport={'width':1440,'height':1000})
                    page=context.new_page();errors=[]
                    page.on('pageerror',lambda error:errors.append(str(error)))
                    page.goto(address)
                    page.wait_for_function('dashboardReady && !document.body.inert')
                    assert page.evaluate("""()=>clients.length===0&&calls.length===0&&feedback.length===0&&backups.length===0&&workspaceRecords().length===0&&invoiceWorkbook===null"""),'A fresh browser must start without sample records'
                    base_record=make_test_record()
                    page.evaluate("""async record=>{clients.push(record);await persist();renderAll()}""",base_record)
                    assert page.locator('#page-dashboard').get_by_role('heading',name='Follow-ups',exact=True).count()==0
                    assert page.locator('#navigationToggle').is_visible()
                    color_options=page.locator('.work-color-select').first.locator('option').all_inner_texts()
                    assert color_options[0]=='—' and 'Follow status' not in color_options and 'Automatic — Untouched' not in color_options,color_options
                    default_color=page.locator('.work-color-select').first.evaluate("""element=>{
                      const selectStyle=getComputedStyle(element),untouched=document.querySelector('.work-legend .work-tag.work-red'),tagStyle=getComputedStyle(untouched),rowStyle=getComputedStyle(document.querySelector('.data-table tr.work-red'));
                      return {selectedValue:element.value,selectColor:selectStyle.color,selectBackground:selectStyle.backgroundColor,tagColor:tagStyle.color,tagBackground:tagStyle.backgroundColor,rowBackground:rowStyle.backgroundColor};
                    }""")
                    assert default_color=={'selectedValue':'','selectColor':'rgb(255, 255, 255)','selectBackground':'rgb(100, 116, 139)','tagColor':'rgb(181, 44, 44)','tagBackground':'rgba(0, 0, 0, 0)','rowBackground':'rgba(0, 0, 0, 0)'},default_color
                    not_interested=page.evaluate("()=>{const node=document.createElement('div');node.innerHTML=badge('Not Interested');document.body.appendChild(node.firstElementChild);const style=getComputedStyle(document.body.lastElementChild);const result={className:document.body.lastElementChild.className,background:style.backgroundColor,color:style.color};document.body.lastElementChild.remove();return result}")
                    assert not_interested=={'className':'status not-interested','background':'rgb(255, 140, 0)','color':'rgb(255, 255, 255)'},not_interested
                    assert page.locator('.work-tag.work-orange').inner_text()=='Orange: Not Interested'
                    assert page.locator('.work-legend .duplicate-neon').inner_text()=='Neon: duplicate records'
                    assert page.locator('.work-legend .duplicate-neon').evaluate("element=>getComputedStyle(element).color")=='rgb(185, 214, 0)'
                    assert page.locator('#newStatus option').all_inner_texts().count('Not Interested')==1
                    assert {'Demand Fulfilled','Order Finalized'}.issubset(set(page.locator('#newStatus option').all_inner_texts()))
                    routing=page.evaluate("""async base=>{
                      const first={...base,id:'HK-DUPTEST1',name:'Neon Duplicate Test',status:'Contacted',workColor:'green'};
                      const second={...first,id:'HK-DUPTEST2',workColor:''};
                      const unique={...first,id:'HK-UNIQUETEST',name:'Single Unique Test'};
                      clients.push(first,unique);sectionRecords.leads.push(second);refreshDuplicateWorkRecords();renderDashboard();
                      const row=document.querySelector('#clientRows tr[data-work-record="HK-DUPTEST1"]');
                      const uniqueRow=document.querySelector('#clientRows tr[data-work-record="HK-UNIQUETEST"]');
                      const nameColor=row&&getComputedStyle(row.querySelector('.client-name')).color;
                      const rowBackground=row&&getComputedStyle(row.cells[0]).backgroundColor;
                      const uniqueIsUnhighlighted=uniqueRow&&!uniqueRow.classList.contains('duplicate-record')&&getComputedStyle(uniqueRow.querySelector('.client-name')).color!=='rgb(185, 214, 0)';
                      await routeGreenPotentialLeads('clients');
                      const copy=sectionRecords.potential.find(record=>record.sourceWorkspace==='clients'&&record.sourceRecordId===first.id);
                      refreshDuplicateWorkRecords();renderPotentialPage();
                      const copyRow=copy&&document.querySelector('#potentialRows tr[data-work-record="'+copy.id+'"]');
                      const result={nameColor,rowBackground,uniqueIsUnhighlighted,sourceRetained:clients.includes(first),copyExists:!!copy,copyIsGreen:copy?.workColor==='green',copyHighlighted:copyRow?.classList.contains('duplicate-record')===true};
                      clients=clients.filter(record=>!['HK-DUPTEST1','HK-UNIQUETEST'].includes(record.id));sectionRecords.leads=sectionRecords.leads.filter(record=>record.id!=='HK-DUPTEST2');sectionRecords.potential=sectionRecords.potential.filter(record=>record.sourceRecordId!=='HK-DUPTEST1');segmentData.potential=sectionRecords.potential;
                      await Promise.all([persist(),persistLeads(),persistSegments()]);renderAll();return result;
                    }""",base_record)
                    assert routing=={'nameColor':'rgb(185, 214, 0)','rowBackground':'rgb(241, 250, 244)','uniqueIsUnhighlighted':True,'sourceRetained':True,'copyExists':True,'copyIsGreen':True,'copyHighlighted':True},routing
                    assert page.locator('#navClients').inner_text()=='1'
                    sections=page.evaluate("Array.from(document.querySelectorAll('.workspace-page'),section=>({id:section.id,title:section.querySelector('h2')?.innerText,parent:section.parentElement.id}))")
                    assert all(section['parent']!='page-dashboard' for section in sections),sections
                    assert page.evaluate("document.querySelector('#clientRows').closest('.page').id==='page-dashboard'"),'The combined records table must stay on Dashboard'
                    sections=[{'id':section['id'],'title':section['title']} for section in sections]
                    assert sections==[{'id':'page-clients','title':'Hospital & Client Business Dashboard'},{'id':'page-rghs','title':'RGHS Data'},{'id':'page-served','title':'Already Served'},{'id':'page-doctors','title':'Doctors'},{'id':'page-potential','title':'HOSPkart Potential Leads'},{'id':'page-leads','title':'Leads & Queries'}],sections
                    page.evaluate("navigate('served')")
                    assert page.locator('#page-served').is_visible() and not page.locator('#page-doctors').is_visible() and not page.locator('#page-dashboard').is_visible()
                    assert page.evaluate("document.querySelectorAll('.page.active').length===1"),'Workspace navigation must activate exactly one page'
                    page.evaluate("navigate('dashboard')")
                    rghs=page.evaluate("""async base=>{
                      const fixture={...base,id:'HK-R-fixture',name:'Independent RGHS fixture',sourceData:{Program:'RGHS'}};
                      const master=JSON.stringify(clients);sectionRecords.rghs.push(fixture);
                      try{await persistRghs();renderRghsPage();renderDashboard();return {count:document.getElementById('rghsTotal').textContent,shown:document.getElementById('rghsRows').innerText.includes(fixture.name),masterUnchanged:master===JSON.stringify(clients),stored:JSON.parse(dashboardStorage.get('hk_v2_rghs')).some(record=>record.id===fixture.id),overallIncludes:document.getElementById('clientRows').innerText.includes(fixture.name)&&document.getElementById('clientRows').innerText.includes('RGHS Data'),overallCount:Number(document.getElementById('kpiClients').textContent),workspaceCount:workspaceRecords().length}}
                      finally{sectionRecords.rghs.pop();await persistRghs();renderRghsPage();renderDashboard()}
                    }""",base_record)
                    assert rghs['count']=='1' and rghs['shown'] and rghs['masterUnchanged'] and rghs['stored'] and rghs['overallIncludes'] and rghs['overallCount']==rghs['workspaceCount'],rghs
                    page.click('button[onclick="loadInvoiceWorkbook(true)"]')
                    page.locator('#toast').get_by_text('No workbook is stored on this server. Use Import Excel to load your file.',exact=True).wait_for()
                    if transfer:
                        page.evaluate("navigate('data')")
                        page.set_input_files('#jsonCheckpointFile',str(transfer))
                        page.locator('#importConfirmAccept').click()
                        page.wait_for_function('invoiceWorkbook?.sheets.length===6')
                        assert page.evaluate('JSON.stringify({clients,calls,feedback,segmentData,sections:sectionRecords,invoiceWorkbook})')==expected_transfer
                    else:
                        page.set_input_files('#excelImportFile',str(fixture))
                        page.wait_for_function('invoiceWorkbook?.sheets.length===6')
                    page.evaluate("navigate('dashboard')")
                    assert page.locator('#page-dashboard').is_visible()
                    assert page.locator('#kpiClients').inner_text().strip()!='0'
                    assert page.get_by_text('All-Workspace Record Headcount',exact=True).is_visible()
                    assert page.locator('#page-dashboard #invoiceDashboard').count()==0
                    assert page.evaluate("Array.from(document.querySelectorAll('.workspace-page')).every(section=>section.parentElement.id!=='page-dashboard'&&!section.classList.contains('hidden'))"),'Workspaces should remain separate pages when an invoice workbook is loaded'
                    page.evaluate("navigate('reports')")
                    assert page.locator('#invoiceDashboard').is_visible()
                    assert page.locator('#invoiceDetailRows tr').count()==3
                    assert '₹600' in page.locator('#invoiceMetrics').inner_text()
                    page.evaluate("navigate('served')")
                    assert page.locator('#servedTotal').inner_text()=='4'
                    assert page.evaluate("segmentData.served.filter(record=>record.name==='Test Alpha').length")==3
                    assert page.locator('#servedOrderValue').inner_text()=='₹600'
                    assert page.locator('#servedDistinctClients').inner_text()=='2'
                    assert page.locator('#servedVisibleCount').inner_text()=='Showing 4 of 4 served entries'
                    page.fill('#servedSearch','Test Alpha')
                    assert page.locator('#servedVisibleCount').inner_text()=='Showing 3 of 4 served entries'
                    assert page.locator('#servedTotal').inner_text()=='4','Search must not change the page-wide headcount'
                    page.fill('#servedSearch','')
                    page.evaluate("navigate('dashboard')")
                    before_filter=page.locator('#kpiClients').inner_text()
                    page.fill('#search','Test Alpha')
                    assert page.locator('#kpiClients').inner_text()==before_filter,'Dashboard filters must not change the all-workspace headcount'
                    assert page.locator('#resultCount').inner_text().startswith('Showing ')
                    page.fill('#search','')
                    assert page.locator('#kpiClients').inner_text()==before_filter
                    page.evaluate("navigate('reports')")
                    for _ in range(20):page.evaluate('installInvoiceWorkbook(invoiceWorkbook,invoiceWorkbook.filename)')
                    assert page.evaluate('backups.length')==10, 'IndexedDB retains all ten checkpoints beyond the localStorage quota'
                    assert page.evaluate('segmentData.served.length')==4, 'Repeated invoice line items must remain separate served records'
                    page.reload();page.wait_for_function('invoiceWorkbook?.sheets.length===6')
                    page.evaluate("navigate('reports')")
                    assert page.locator('#invoiceDetailRows tr').count()==3
                    # A write failure must roll back active data and its checkpoint.
                    rollback=page.evaluate("""async()=>{
                      const keys=['hk_v2_clients','hk_v2_calls','hk_v2_feedback','hk_v2_segments','hk_v2_invoice_workbook','hk_v2_backups'];
                      const before=JSON.stringify(keys.map(key=>dashboardStorage.get(key)));
                      const beforeState=JSON.stringify({clients,calls,feedback,segmentData,invoiceWorkbook,backups});
                      const original=IDBObjectStore.prototype.put;let failed=false,message='';
                      IDBObjectStore.prototype.put=function(value,key){if(key==='hk_v2_segments'&&!failed){failed=true;throw new DOMException('Test quota','QuotaExceededError')}return original.call(this,value,key)};
                      try{await installInvoiceWorkbook(invoiceWorkbook,'failed-import.xlsx')}catch(error){message=error.message}finally{IDBObjectStore.prototype.put=original}
                      const stored=await databaseRecords();
                      return {message,same:before===JSON.stringify(keys.map(key=>dashboardStorage.get(key))),sameDisk:before===JSON.stringify(keys.map(key=>stored.get(key))),sameState:beforeState===JSON.stringify({clients,calls,feedback,segmentData,invoiceWorkbook,backups})};
                    }""")
                    assert rollback['same'] and rollback['sameDisk'] and rollback['sameState'] and rollback['message'],rollback
                    # A normal form waits for durability, and a failed form does
                    # not leave an unsaved edit in memory or report success.
                    page.evaluate("""async()=>{
                      openWorkStatusEditor(clients[0].id);
                      document.getElementById('workRecordRemark').value='Persisted work remark';
                      await saveWorkStatus({preventDefault(){}});
                    }""")
                    assert page.evaluate('clients[0].remark')=='Persisted work remark'
                    failed_form=page.evaluate("""async()=>{
                      const before=JSON.stringify(clients),original=IDBObjectStore.prototype.put;
                      openWorkStatusEditor(clients[0].id);
                      document.getElementById('workRecordRemark').value='Must not be saved';
                      IDBObjectStore.prototype.put=function(){throw new DOMException('Form save failed','QuotaExceededError')};
                      try{await saveWorkStatus({preventDefault(){}})}finally{IDBObjectStore.prototype.put=original;closeDrawer('workStatusDrawer')}
                      return before===JSON.stringify(clients);
                    }""")
                    assert failed_form
                    assert 'Changes were not saved' in page.locator('#toast').inner_text()
                    page.reload();page.wait_for_function('invoiceWorkbook?.sheets.length===6')
                    page.evaluate("navigate('dashboard')")
                    assert page.locator('#page-dashboard').is_visible()
                    assert page.locator('#kpiClients').inner_text().strip()!='0'
                    page.evaluate("navigate('reports')")
                    assert page.locator('#invoiceDashboard').is_visible()
                    assert page.evaluate('clients[0].remark')=='Persisted work remark'
                    page.evaluate("navigate('data')")
                    with page.expect_download() as download:page.click('button[onclick="prepareFullBackup()"]')
                    transfer=root/'transfer.json';download.value.save_as(transfer)
                    expected_transfer=page.evaluate('JSON.stringify({clients,calls,feedback,segmentData,sections:sectionRecords,invoiceWorkbook})')
                    page.set_viewport_size({'width':390,'height':844})
                    assert page.locator('#navigationToggle').is_visible()
                    page.click('#navigationToggle')
                    assert page.locator('#mainNavigation').get_by_role('button',name='✓ Already Served').is_visible()
                    page.locator('#mainNavigation .nav-item[data-view="served"]').click()
                    page.wait_for_function("currentView==='served' && document.getElementById('page-served').classList.contains('active')")
                    served_state=page.locator('#page-served').evaluate("element=>({display:getComputedStyle(element).display,visibility:getComputedStyle(element).visibility,height:element.getBoundingClientRect().height,dashboardDisplay:getComputedStyle(document.getElementById('page-dashboard')).display})")
                    assert served_state['display']!='none' and served_state['visibility']=='visible' and served_state['height']>0 and served_state['dashboardDisplay']=='none',served_state
                    page.click('#navigationToggle')
                    page.locator('#mainNavigation .nav-item[data-view="data"]').click()
                    assert page.locator('#page-data').is_visible()
                    assert page.evaluate('document.documentElement.scrollWidth<=window.innerWidth')
                    row=page.locator('#importRemovalRows tr').filter(has=page.locator('td').get_by_text('Categorized Items',exact=True))
                    row.get_by_role('button',name='Remove',exact=True).click()
                    page.fill('#importRemovalPassword',password);page.fill('#importRemovalConfirmation','REMOVE Categorized Items');page.click('#importRemovalSubmit')
                    page.wait_for_function('invoiceWorkbook.sheets.length===5')
                    page.evaluate("navigate('reports')")
                    assert '₹600' in page.locator('#invoiceMetrics').inner_text()
                    assert not errors,errors
                    context.close()
                    # Reproduce the screenshot: 1,743 existing clients and a full
                    # legacy localStorage must migrate without dropping history.
                    large=browser.new_context();large.add_init_script(f"""(() => {{
                      const base={json.dumps(base_record)};
                      const records=Array.from({{length:1743}},(_,i)=>({{...base,id:'legacy-'+i,name:'Existing Hospital '+i,sourceData:{{originalColumn:'preserved-'+i,details:'x'.repeat(200)}}}}));
                      records[0].sourceData.Program='RGHS';
                      const note={{clientId:'legacy-0',date:'2026-09-01',purpose:'Existing call',connected:'Yes',notes:'Must survive migration'}};
                      const savedFeedback={{clientId:'legacy-0',date:'2026-09-01',rating:'5',text:'Existing feedback'}};
                      const segments={{served:[],doctors:[{{...base,id:'old-doctor',name:'Existing Doctor'}}],potential:[]}};
                      const checkpoint={{id:'legacy-backup',createdAt:new Date().toISOString(),reason:'Existing checkpoint',clients:records,calls:[note],feedback:[savedFeedback],segments,invoiceWorkbook:null}};
                      const data={{hk_v2_clients:records,hk_v2_calls:[note],hk_v2_feedback:[savedFeedback],hk_v2_segments:segments,hk_v2_backups:[checkpoint]}};
                      Object.entries(data).forEach(([key,value])=>localStorage.setItem(key,JSON.stringify(value)));
                      let chunks=0;try{{while(chunks<100){{localStorage.setItem('unrelated-'+chunks,'z'.repeat(100000));chunks++;}}}}catch(error){{if(error.name!=='QuotaExceededError')throw error;}}
                    }})()""")
                    p=large.new_page();p.goto(address);p.wait_for_function('dashboardReady && !document.body.inert')
                    seed=p.evaluate("""()=>({records:JSON.stringify(clients),calls:JSON.stringify(calls),feedback:JSON.stringify(feedback),chunks:Array.from({length:100},(_,i)=>localStorage.getItem('unrelated-'+i)).filter(Boolean).length})""")
                    assert seed['chunks']>0
                    assert p.locator('#navClients').inner_text()=='1743'
                    assert p.evaluate('JSON.stringify(clients)')==seed['records']
                    assert p.evaluate("localStorage.getItem('hk_v2_clients')") is None
                    assert p.evaluate("localStorage.getItem('unrelated-0')")=='z'*100000
                    p.set_input_files('#excelImportFile',str(fixture));p.wait_for_function('segmentData.served.length===4')
                    assert p.evaluate('clients.length')==1745
                    assert p.evaluate("JSON.stringify(clients.filter(client=>client.id.startsWith('legacy-')))")==seed['records']
                    assert p.evaluate('JSON.stringify(calls)')==seed['calls']
                    assert p.evaluate('JSON.stringify(feedback)')==seed['feedback']
                    assert p.evaluate('segmentData.doctors.length')==1
                    assert p.evaluate("sectionRecords.doctors.length===1 && JSON.parse(dashboardStorage.get('hk_v2_doctors')).length===1")
                    assert p.evaluate("sectionRecords.leads.every(record=>record.id.startsWith('HK-L')) && sectionRecords.leads.length===1743")
                    separation=p.evaluate("""async()=>{
                      const clientRemark=clients[0].remark,lead=sectionRecords.leads[0],rghs=sectionRecords.rghs[0];
                      lead.remark='Independent lead edit';rghs.remark='Independent RGHS edit';
                      await Promise.all([persistLeads(),persistRghs()]);
                      return {masterUnchanged:clients[0].remark===clientRemark,leadStored:JSON.parse(dashboardStorage.get('hk_v2_leads'))[0].remark,rghsStored:JSON.parse(dashboardStorage.get('hk_v2_rghs'))[0].remark,rghsId:rghs.id};
                    }""")
                    assert separation=={'masterUnchanged':True,'leadStored':'Independent lead edit','rghsStored':'Independent RGHS edit','rghsId':'HK-R0001'},separation
                    assert p.evaluate("JSON.stringify(clients)===JSON.stringify(JSON.parse(dashboardStorage.get('hk_v2_clients')))")
                    assert p.evaluate('backups[0].clients.length')==1743
                    p.reload();p.wait_for_function('invoiceWorkbook?.sheets.length===6')
                    assert p.evaluate('clients.length')==1745
                    assert p.evaluate('segmentData.served.length')==4
                    assert p.evaluate("sectionRecords.leads[0].remark==='Independent lead edit' && sectionRecords.rghs[0].remark==='Independent RGHS edit' && clients[0].remark!=='Independent lead edit'")
                    # A second tab sees committed changes, rather than overwriting
                    # the first tab with an obsolete dataset after a reset.
                    peer=large.new_page();peer.goto(address);peer.wait_for_function('dashboardReady && !document.body.inert')
                    p.evaluate("navigate('data')")
                    p.check('#resetAcknowledge');p.fill('#resetPhrase','RESET DASHBOARD')
                    p.click('#resetConfirmButton')
                    p.wait_for_function('!dashboardResetting && clients.length===0 && invoiceWorkbook===null && Object.values(sectionRecords).every(records=>records.length===0)')
                    peer.wait_for_function('dashboardReady && clients.length===0 && invoiceWorkbook===null && Object.values(sectionRecords).every(records=>records.length===0)')
                    p.reload();p.wait_for_function('dashboardReady && !document.body.inert')
                    assert p.evaluate("clients.length===0&&calls.length===0&&feedback.length===0&&invoiceWorkbook===null&&Object.values(sectionRecords).every(records=>records.length===0)")
                    assert not errors,errors
                    large.close()
                    failed_migration=browser.new_context()
                    failed_migration.add_init_script("""
                      localStorage.setItem('hk_v2_clients',JSON.stringify([{id:'old-record',name:'Preserved Hospital',city:'Test',status:'New Lead',quote:0,order:0,next:'',remark:''}]));
                      const put=IDBObjectStore.prototype.put;
                      IDBObjectStore.prototype.put=function(){throw new DOMException('Migration failed','QuotaExceededError')};
                    """)
                    p=failed_migration.new_page();p.goto(address)
                    p.wait_for_function('!document.body.inert')
                    p.locator('#browserStorageNotice:not(.hidden)').wait_for()
                    assert p.evaluate("JSON.parse(localStorage.getItem('hk_v2_clients'))[0].name")=='Preserved Hospital'
                    assert p.evaluate('clients[0].name')=='Preserved Hospital'
                    p.evaluate("navigate('data')")
                    with p.expect_download() as emergency:p.click('button[onclick="prepareFullBackup()"]')
                    emergency.value.save_as(root/'emergency.json')
                    assert json.loads((root/'emergency.json').read_text(encoding='utf-8'))['clients'][0]['name']=='Preserved Hospital'
                    assert p.evaluate('async()=> (await databaseRecords()).size')==0
                    failed_migration.close()
                    # Old/corrupted state must not crash initialization or be erased.
                    for key,value in [('hk_v2_clients','{}'),('hk_v2_clients','null'),('hk_v2_clients','[null]'),('hk_v2_backups','[{}]'),('hk_v2_segments','{"served":[null]}'),('hk_v2_invoice_workbook','{"sheets":null}')]:
                        old=browser.new_context();old.add_init_script(f'localStorage.setItem({json.dumps(key)},{json.dumps(value)})')
                        p=old.new_page();failures=[];p.on('pageerror',lambda error:failures.append(str(error)));p.goto(address)
                        p.locator('#browserStorageNotice:not(.hidden)').wait_for()
                        assert not failures,(key,failures)
                        assert p.evaluate(f'localStorage.getItem({json.dumps(key)})')==value
                        old.close()
                    blocked=browser.new_context();blocked.add_init_script("Object.defineProperty(window,'localStorage',{get(){throw new DOMException('Storage blocked','SecurityError')}})")
                    p=blocked.new_page();failures=[];p.on('pageerror',lambda error:failures.append(str(error)));p.goto(address)
                    p.locator('#browserStorageNotice:not(.hidden)').wait_for();assert not failures,failures
                    blocked.close();reports.append({'browser':channel,'passed':True,'repeated_imports':20,'corrupt_storage_cases':6,'blocked_storage':True,'transfer':True,'mobile':True,'admin_removal':True,'rollback':True,'full_legacy_storage_clients':1743,'migration_failure_preserves_originals':True,'emergency_export':True,'cross_tab_reset':True})
                    browser.close()
            print(json.dumps({'server':'Waitress','fresh_setup_without_private_data':True,'results':reports}))
        finally:
            server.close()
            if app.scheduler.running:app.scheduler.shutdown(wait=False)
            if previous_hash is not None:os.environ['ADMIN_PASSWORD_HASH']=previous_hash


if __name__=='__main__':check()
