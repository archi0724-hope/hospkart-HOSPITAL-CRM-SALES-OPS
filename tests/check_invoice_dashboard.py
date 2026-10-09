"""End-to-end invoice dashboard checks using an isolated synthetic workbook."""
import os
import sys
import tempfile
from pathlib import Path
from threading import Thread

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app
from openpyxl import Workbook
from playwright.sync_api import sync_playwright
from werkzeug.serving import make_server


def make_workbook(path):
    workbook=Workbook();workbook.remove(workbook.active)
    invoices=[['2026-09-01','INV-1','Test Alpha','Jaipur',100,'Paid'],['2026-09-02','INV-2','Test Alpha','Jaipur',200,'Active'],['2026-09-03','INV-3','Test Beta','Delhi',300,'Part Paid']]
    headers=['Invoice Date','Invoice No.','Client / Hospital','Delivery City / Location','Invoice Total (₹)','Invoice Status']
    summary=workbook.create_sheet('Client Summary');summary.append(headers)
    for row in invoices:summary.append(row)
    details=workbook.create_sheet('Client Call List');details.append(headers+['Item Name','Category','Item Total incl. GST (₹)'])
    for row,item,amount in [(invoices[0],'Item A',40),(invoices[0],'Item B',60),(invoices[1],'Item C',200),(invoices[2],'Item D',300)]:
        details.append(row+[item,'Test category',amount])
    cancelled=workbook.create_sheet('Cancelled Invoices');cancelled.append(headers)
    cancelled.append(['2026-09-04','INV-X','Test Cancelled','Delhi',50,'Canceled'])
    for name in ['Categorized Items','Verification Summary','Read Me']:
        sheet=workbook.create_sheet(name);sheet.append(['Label','Value']);sheet.append(['Synthetic test source',name])
    workbook.save(path);workbook.close()


def check():
    with tempfile.TemporaryDirectory() as temp:
        app.DATA_DIR = Path(temp)
        app.DB_PATH = app.DATA_DIR / 'test.db'
        app.init_db()
        invoice_file=app.DATA_DIR / 'synthetic-invoice-fixture.xlsx'
        make_workbook(invoice_file)
        server = make_server('127.0.0.1', 5011, app.app)
        thread = Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(channel=os.getenv('BROWSER_CHANNEL','chrome'))
                page = browser.new_page(viewport={'width': 1440, 'height': 1000})
                errors = []
                page.on('pageerror', lambda error: errors.append(str(error)))
                page.goto('http://127.0.0.1:5011')
                page.wait_for_function('dashboardReady && !document.body.inert')
                page.set_input_files('#excelImportFile', str(invoice_file))
                page.wait_for_function('invoiceWorkbook?.sheets.length===6')
                assert page.locator('#page-dashboard').is_visible()
                assert page.locator('#invoiceDashboard').count() == 1
                assert page.locator('#page-dashboard #invoiceDashboard').count() == 0
                assert page.get_by_text('All-Workspace Record Headcount', exact=True).is_visible()
                assert page.locator('#kpiClients').inner_text().strip() != '0'
                page.evaluate("navigate('reports')")
                assert page.locator('#invoiceDashboard').is_visible()
                page.locator('#invoiceMetrics').get_by_text('2', exact=True).wait_for()
                metrics = page.locator('#invoiceMetrics').inner_text()
                assert '3' in metrics and '4' in metrics and '6' in metrics, metrics
                assert page.locator('#invoiceDetailRows tr').count() == 3
                assert page.locator('#servedTotal').inner_text() == '4'
                assert page.evaluate("segmentData.served.length") == 4
                assert page.evaluate("segmentData.served.filter(record=>record.name==='Test Alpha').length") > 1
                actual = page.evaluate("invoiceSheet('Client Summary').reduce((sum,row)=>sum+row['Invoice Total (₹)'],0)")
                assert round(actual, 2) == 600, actual
                page.select_option('#invoiceDetailSheet', 'Client Call List')
                assert page.locator('#invoiceDetailRows tr').count() == 4
                assert page.locator('#invoiceDetailHead th').count() == 9
                page.select_option('#invoiceDetailSheet', 'Cancelled Invoices')
                assert page.locator('#invoiceDetailRows tr').count() == 1
                page.select_option('#invoiceStatus', 'Paid')
                assert page.evaluate("uniqueInvoices(invoiceSheet('Client Summary')).filter(row=>row['Invoice Status']==='Paid').length") == 1
                assert '1' in page.locator('#invoiceMetrics').inner_text()
                page.click('button[onclick="clearInvoiceFilters()"]')
                page.select_option('#invoiceDetailSheet', 'Client Summary')
                with page.expect_download() as download:
                    page.click('button[onclick="exportInvoiceDetails()"]')
                assert download.value.suggested_filename.endswith('.xls')
                page.fill('#invoiceSearch', 'nothing-matches-this-query')
                assert 'No matching records.' in page.locator('#invoiceDetailRows').inner_text()
                assert '₹0' not in page.locator('#invoiceMetrics').inner_text()
                page.click('button[onclick="clearInvoiceFilters()"]')
                page.reload()
                page.evaluate("navigate('reports')")
                page.locator('#invoiceMetrics').get_by_text('2', exact=True).wait_for()
                assert page.locator('#servedTotal').inner_text() == '4'
                # Older saved workbooks had grouped client records. The saved
                # workbook must rebuild the item-level served list on startup.
                page.evaluate("""async()=>{
                  const legacy=segmentData.served.slice(0,38).map(record=>({...record}));
                  legacy[0].workColor='yellow';legacy[0].remark='Preserve existing served remark';
                  await writeDashboardValues({hk_v2_segments:{...segmentData,served:legacy},hk_v2_served:legacy});
                }""")
                page.reload()
                page.wait_for_function("invoiceWorkbook && segmentData.served.length===4")
                assert page.locator('#servedTotal').inner_text() == '4'
                assert page.evaluate("segmentData.served[0].workColor==='yellow'&&segmentData.served[0].remark==='Preserve existing served remark'")
                output = Path(temp) / 'invoice-dashboard-desktop.png'
                page.screenshot(path=str(output), full_page=True)
                page.set_viewport_size({'width': 390, 'height': 844})
                page.screenshot(path=str(output.with_name('invoice-dashboard-mobile.png')), full_page=True)
                assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth'), 'Mobile page overflows'
                # Re-import must retain all rows and not duplicate client master records.
                before = page.evaluate('clients.length')
                file = Path(temp) / 'invoice-reimport.xlsx'
                make_workbook(file)
                page.set_input_files('#servedImportFile', str(file))
                page.wait_for_function("invoiceWorkbook.filename === 'invoice-reimport.xlsx'")
                assert page.evaluate('clients.length') == before
                assert page.locator('#servedTotal').inner_text() == '4'
                page.evaluate("navigate('dashboard')")
                assert page.locator('#page-dashboard').is_visible()
                assert page.locator('#kpiClients').inner_text().strip() != '0'
                assert page.locator('#invoiceDashboard').is_hidden()
                # Reset suppression must persist across reloads.
                page.evaluate("writeDashboardValues({hk_v2_invoice_workbook:null,hk_v2_reset_revision:'test'})")
                page.reload()
                page.wait_for_timeout(500)
                assert page.locator('#invoiceDashboard').is_hidden()
                assert not errors, errors
                browser.close()
                print({'passed': True, 'clients': 2, 'invoices': 3, 'item_rows': 4, 'cancelled_item_rows': 1, 'invoice_value': actual, 'browser_errors': errors})
        finally:
            server.shutdown()
            app.scheduler.shutdown(wait=False)


if __name__ == '__main__':
    check()
