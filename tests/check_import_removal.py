"""Verify scoped import removal and restore with an isolated synthetic workbook."""
import sys
import tempfile
from pathlib import Path
from threading import Thread

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app
from openpyxl import Workbook
from playwright.sync_api import sync_playwright
from werkzeug.security import generate_password_hash
from werkzeug.serving import make_server


def make_workbook(path):
    workbook=Workbook();workbook.remove(workbook.active)
    invoices=[['2026-09-01','INV-1','Test Alpha','Jaipur',100,'Paid'],['2026-09-02','INV-2','Test Alpha','Jaipur',200,'Active'],['2026-09-03','INV-3','Test Beta','Delhi',300,'Part Paid']]
    headers=['Invoice Date','Invoice No.','Client / Hospital','Delivery City / Location','Invoice Total (₹)','Invoice Status']
    summary=workbook.create_sheet('Client Summary');summary.append(headers)
    for row in invoices:summary.append(row)
    details=workbook.create_sheet('Client Call List');details.append(headers+['Item Name','Category','Item Total incl. GST (₹)'])
    for row,item,amount in [(invoices[0],'Item A',40),(invoices[0],'Item B',60),(invoices[1],'Item C',200),(invoices[2],'Item D',300)]:
        details.append(row+[item,'Test category',amount])
    cancelled=workbook.create_sheet('Cancelled Invoices');cancelled.append(headers)
    cancelled.append(['2026-09-04','INV-X','Test Cancelled','Delhi',50,'Canceled'])
    for name in ['Categorized Items','Verification Summary','Read Me']:
        sheet=workbook.create_sheet(name);sheet.append(['Label','Value']);sheet.append(['Synthetic test source',name])
    workbook.save(path);workbook.close()


def check():
    with tempfile.TemporaryDirectory() as temp:
        app.DATA_DIR = Path(temp)
        app.DB_PATH = app.DATA_DIR / 'test.db'
        app.ADMIN_PASSWORD_PATH = app.DATA_DIR / 'admin.hash'
        app.init_db()
        app.admin_attempts.clear()
        password = 'browser-test-admin-only'
        fixture=app.DATA_DIR / 'synthetic-invoice-fixture.xlsx'
        make_workbook(fixture)
        server = make_server('127.0.0.1', 5012, app.app)
        Thread(target=server.serve_forever, daemon=True).start()
        try:
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(channel='msedge')
                page = browser.new_page(viewport={'width':1440,'height':1000})
                errors=[]
                page.on('pageerror',lambda error:errors.append(str(error)))
                page.goto('http://127.0.0.1:5012')
                page.wait_for_function('dashboardReady && !document.body.inert')
                page.set_input_files('#excelImportFile',str(fixture))
                page.wait_for_function('invoiceWorkbook?.sheets.length===6')
                page.evaluate("navigate('data')")
                page.screenshot(path=str(Path(temp) / 'import-removal-desktop.png'),full_page=True)
                before=page.evaluate('JSON.stringify({clients,calls,feedback,segmentData})')
                original=page.evaluate('JSON.stringify(invoiceWorkbook)')

                def open_remove(name):
                    page.locator('#importRemovalRows tr').filter(has=page.locator('td').get_by_text(name,exact=True)).get_by_role('button',name='Remove',exact=True).click()

                def submit(name, secret=password):
                    page.fill('#importRemovalPassword',secret)
                    page.fill('#importRemovalConfirmation','REMOVE '+name)
                    page.click('#importRemovalSubmit')

                open_remove('Categorized Items')
                page.screenshot(path=str(Path(temp) / 'import-removal-dialog.png'))
                submit('Categorized Items')
                page.get_by_text('An administrator must configure the admin password before removing imports.',exact=True).wait_for()
                assert page.evaluate('JSON.stringify(invoiceWorkbook)')==original
                app.ADMIN_PASSWORD_PATH.write_text(generate_password_hash(password),encoding='utf-8')
                submit('Categorized Items','wrong')
                page.get_by_text('Incorrect admin password. No data was removed.',exact=True).wait_for()
                assert page.evaluate('JSON.stringify(invoiceWorkbook)')==original
                submit('Categorized Items')
                page.wait_for_function('invoiceWorkbook.sheets.length===5')
                assert page.evaluate('JSON.stringify({clients,calls,feedback,segmentData})')==before
                assert page.evaluate("!invoiceWorkbook.sheet_previews['Categorized Items']")
                assert page.evaluate('backups[0].invoiceWorkbook.sheets.length')==6
                page.reload()
                page.wait_for_function('invoiceWorkbook?.sheets.length===5')
                assert not page.locator('#invoiceDetailSheet option').get_by_text('Categorized Items',exact=True).count()
                # Intentional restore brings the removed worksheet back.
                page.on('dialog',lambda dialog:dialog.accept())
                page.evaluate('restoreBackup(backups[0].id)')
                page.wait_for_function('invoiceWorkbook.sheets.length===6')
                page.evaluate("navigate('data')")
                open_remove('Client Summary'); submit('Client Summary')
                page.wait_for_function('invoiceWorkbook.sheets.length===5')
                page.evaluate("navigate('reports')")
                assert page.locator('#invoiceMetrics').get_by_text('3',exact=True).count()==1
                assert page.locator('#invoiceMetrics').get_by_text('4',exact=True).count()==1
                assert 'Client Summary removed' in page.locator('#invoiceCoverage').inner_text()
                page.evaluate("navigate('data')")
                open_remove('Client Call List'); submit('Client Call List')
                page.wait_for_function('invoiceWorkbook.sheets.length===4')
                page.evaluate("navigate('reports')")
                assert page.locator('#invoiceMetrics').get_by_text('Unavailable',exact=True).count()==5
                assert page.evaluate('JSON.stringify({clients,calls,feedback,segmentData})')==before
                # A section removal clears that section only.
                page.evaluate("async()=>{sectionRecords.doctors=[{id:'test-doctor',name:'Test doctor',sourceData:{Name:'Test doctor'}}];segmentData.doctors=sectionRecords.doctors;await persistSegments();renderAll()}")
                served=page.evaluate('JSON.stringify(segmentData.served)')
                page.evaluate("navigate('data')")
                open_remove('Doctors'); submit('doctors')
                page.wait_for_function('segmentData.doctors.length===0')
                assert page.evaluate('JSON.stringify(segmentData.served)')==served
                filename=page.evaluate('invoiceWorkbook.filename')
                open_remove(filename); submit(filename)
                page.wait_for_function('invoiceWorkbook===null')
                assert page.evaluate('JSON.stringify(segmentData.served)')==served
                page.reload();page.wait_for_timeout(300)
                assert page.locator('#invoiceDashboard').is_hidden()
                assert not errors,errors
                browser.close()
                print('Admin denial, exact worksheet removal, other-data preservation, reload, restore, source fallback, section removal, and workbook removal passed. No browser errors.')
        finally:
            server.shutdown()
            app.scheduler.shutdown(wait=False)


if __name__=='__main__':
    check()
