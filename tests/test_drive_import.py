"""Shared Drive import links, private downloads and failures use isolated fixtures."""
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

from crm_test_environment import app
from drive_import import DriveWorkbook, XLSX
from spreadsheet_sync import SheetError

LINK='https://drive.google.com/file/d/test_private_workbook_123456/view'
HEADERS={'X-CRM-Request':'1'}


class DriveImportTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.db=patch.object(app,'DB_PATH',Path(self.temp.name)/'crm.db');self.db.start();self.addCleanup(self.db.stop)
        app.init_db();self.client=app.app.test_client();self.peer=app.app.test_client()

    def post(self,path,payload):
        return self.client.post('/api/drive-import/'+path,json=payload,headers=HEADERS)

    def save(self):
        result=self.post('save',{'link':LINK,'label':'Private hospitals','target':'clients'})
        self.assertEqual(result.status_code,200,result.json)
        return result.json['links'][0]

    def test_link_shared_persistent_and_repeat_save_does_not_duplicate(self):
        link=self.save();self.save();app.init_db()
        state=self.peer.get('/api/drive-import').json
        self.assertEqual(len(state['links']),1);self.assertEqual(state['links'][0]['id'],link['id'])
        self.assertEqual(state['links'][0]['link'],LINK)
        self.assertEqual(self.client.get('/api/workspace').json['workspace']['values']['hk_v2_clients'],[])

    def test_invalid_links_targets_and_labels_never_save(self):
        for changes in [{'link':'http://localhost/private'},{'link':'https://docs.google.com.evil.test/spreadsheets/d/test_private_workbook_123456'},
                        {'link':'javascript:alert(1)'},{'target':[]},{'target':'credentials'},{'label':''},{'label':'x'*121}]:
            payload={'link':LINK,'label':'Private workbook','target':'clients',**changes}
            self.assertEqual(self.post('save',payload).status_code,400)
        self.assertEqual(self.client.get('/api/drive-import').json['links'],[])

    def test_remove_link_keeps_records_and_drive_source_unchanged(self):
        item=self.save()
        response=self.post('remove',{'id':item['id']})
        self.assertEqual(response.json['links'],[])
        self.assertEqual(self.post('download',{'id':item['id']}).status_code,404)

    def test_origin_and_json_headers_required(self):
        self.assertEqual(self.client.post('/api/drive-import/save',json={}).status_code,403)
        response=self.client.post('/api/drive-import/save',json={'link':LINK},headers={**HEADERS,'Origin':'https://evil.example'})
        self.assertEqual(response.status_code,403)

    def test_private_download_is_separate_from_import_and_retains_failures(self):
        item=self.save()
        with patch.object(app.drive_import_provider,'download',return_value=(BytesIO(b'fixture'),'test.xlsx','.xlsx')):
            response=self.post('download',{'id':item['id']})
            self.assertEqual(response.status_code,200);self.assertEqual(response.data,b'fixture')
            self.assertIn('no-store',response.headers['Cache-Control'])
            response.close()
        with patch.object(app.drive_import_provider,'download',side_effect=SheetError('Missing Viewer permission.',403)):
            self.assertEqual(self.post('download',{'id':item['id']}).status_code,403)
        self.assertEqual(len(self.client.get('/api/drive-import').json['links']),1)
        self.assertEqual(self.client.get('/api/workspace').json['workspace']['values']['hk_v2_clients'],[])

    def test_resource_key_kept_and_malicious_query_parameters_removed(self):
        response=self.post('save',{'link':LINK+'?resourcekey=safe-key_123&redirect=https://evil.test','label':'Resource link','target':'served'})
        self.assertEqual(response.json['links'][0]['link'],LINK+'?resourcekey=safe-key_123')

    def test_saved_workbook_copy_is_persistent_and_explicit_not_a_live_drive_read(self):
        from openpyxl import Workbook
        item=self.save();book=Workbook();book.active.append(['Name']);book.active.append(['Synthetic Hospital'])
        output=BytesIO();book.save(output);content=output.getvalue()
        result=self.client.post('/api/drive-import/save-copy',data={'id':item['id'],'file':(BytesIO(content),'copy.xlsx')},headers=HEADERS)
        self.assertEqual(result.status_code,200,result.json)
        self.assertTrue(result.json['links'][0]['copy_saved_at'])
        app.init_db()
        with patch.object(app.drive_import_provider,'download') as google:
            response=self.post('download',{'id':item['id'],'saved_copy':True})
            self.assertEqual(response.data,content);response.close();google.assert_not_called()
        invalid=self.client.post('/api/drive-import/save-copy',data={'id':item['id'],'file':(BytesIO(b'bad'),'copy.xlsx')},headers=HEADERS)
        self.assertEqual(invalid.status_code,400)
        response=self.post('download',{'id':item['id'],'saved_copy':True});self.assertEqual(response.data,content);response.close()
        self.assertEqual(self.client.get('/api/workspace').json['workspace']['values']['hk_v2_clients'],[])

    def test_native_export_and_excel_download_use_fixed_google_endpoints(self):
        provider=DriveWorkbook();session=MagicMock();session.__enter__.return_value=session
        download=MagicMock();download.__enter__.return_value=download;download.status_code=200;download.iter_content.return_value=[b'xlsx-content']
        session.get.return_value=download
        for mime,native in [('application/vnd.google-apps.spreadsheet',True),(XLSX,False),('application/vnd.ms-excel',False)]:
            with patch.object(provider,'session',return_value=session),patch.object(provider,'get',return_value={'mimeType':mime,'name':'Workbook','capabilities':{'canDownload':True}}):
                content,name,suffix=provider.download('test_private_workbook_123456')
            self.assertEqual(content.read(),b'xlsx-content')
            args,kwargs=session.get.call_args
            self.assertTrue(args[0].startswith('https://www.googleapis.com/drive/v3/files/test_private_workbook_123456'))
            self.assertEqual(args[0].endswith('/export'),native)
            self.assertFalse(kwargs['allow_redirects'])

    def test_unsupported_and_disallowed_files_and_oversize_streams_rejected(self):
        provider=DriveWorkbook();session=MagicMock();session.__enter__.return_value=session
        for meta in [{'mimeType':'application/pdf'},{'mimeType':XLSX,'capabilities':{'canDownload':False}}]:
            with patch.object(provider,'session',return_value=session),patch.object(provider,'get',return_value=meta),self.assertRaises(SheetError):
                provider.download('test_private_workbook_123456')
        download=MagicMock();download.__enter__.return_value=download;download.status_code=200
        download.iter_content.return_value=[b'abcd',b'efgh'];session.get.return_value=download
        with patch.object(provider,'session',return_value=session),patch.object(provider,'get',return_value={'mimeType':XLSX,'name':'Workbook'}),patch('drive_import.MAX_BYTES',5),self.assertRaises(SheetError) as caught:
            provider.download('test_private_workbook_123456')
        self.assertEqual(caught.exception.status,413)
