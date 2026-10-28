"""Isolated SQLite/API tests. Google calls are mocked; no live CRM data changes."""
import copy
import os
import sqlite3
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from flask.testing import FlaskClient
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from crm_test_environment import app
from spreadsheet_sync import GoogleSheets, SheetError, parse_link


LINK = 'https://docs.google.com/spreadsheets/d/private_sheet_123456/edit#gid=0'
HEADERS = {'X-CRM-Request': '1'}


class FakeSheets:
    def __init__(self):
        self.rows = [
            [['Lead ID', 'Name', 'Phone', 'Status', 'Created Date', 'Follow-up Date'],
             ['0001', 'Hospital A', '0012345678', 'Query Received', '2026-10-01', '2026-11-01']],
            [['Query ID', 'Lead ID', 'Subject', 'Details', 'Updated'],
             ['0007', '0001', 'Monitor', 'Need a quote', '2026-10-02']],
        ]
        self.failure = None

    def inspect(self, file_id):
        if self.failure:
            raise self.failure
        return {'id': file_id, 'name': 'Private Sales', 'url': LINK.split('#')[0],
                'tabs': [{'id': 1, 'name': 'Sales'}, {'id': 2, 'name': 'Enquiries'}]}

    def read(self, file_id, tabs, headers_only=False):
        if self.failure:
            raise self.failure
        return copy.deepcopy(self.rows)


def connection_payload():
    return {'link': LINK, 'tabs': {'leads': 'Sales', 'queries': 'Enquiries'},
            'auto_minutes': 0, 'mappings': {
                'leads': {'id': 'Lead ID', 'name': 'Name', 'mobile': 'Phone', 'status': 'Status',
                          'created_date': 'Created Date', 'next': 'Follow-up Date'},
                'queries': {'id': 'Query ID', 'lead_id': 'Lead ID', 'subject': 'Subject',
                            'details': 'Details', 'updated_date': 'Updated'}}}


class SpreadsheetTests(unittest.TestCase):
    def setUp(self):
        policy = patch.dict(app.app.config, {'ADMIN_ACCESS_REQUIRED': True})
        policy.start(); self.addCleanup(policy.stop)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.db_patch = patch.object(app, 'DB_PATH', Path(self.temp.name) / 'crm.db')
        self.db_patch.start()
        self.addCleanup(self.db_patch.stop)
        app.init_db()
        self.provider = FakeSheets()
        self.provider_patch = patch.object(app.sheet_sync, 'provider', self.provider)
        self.provider_patch.start()
        self.addCleanup(self.provider_patch.stop)
        admin_patch = patch.object(app, 'admin_password_hash', return_value=generate_password_hash('sheet-test-admin'))
        admin_patch.start()
        self.addCleanup(admin_patch.stop)
        app.admin_attempts.clear()
        self.client = app.app.test_client()

    def post(self, path, payload=None):
        return self.client.post('/api/spreadsheet/' + path, json={**(payload or {}), 'admin_password': 'sheet-test-admin'}, headers=HEADERS)

    def status(self):
        return self.client.get('/api/spreadsheet/status').json

    def connect(self, payload=None):
        preview = self.post('preview', payload or connection_payload())
        self.assertEqual(preview.status_code, 200, preview.json)
        result = self.post('import', {'token': preview.json['token']})
        self.assertEqual(result.status_code, 200, result.json)
        return result.json

    def test_settings_require_admin_even_with_direct_api_calls(self):
        original = self.status()
        for path in ('source', 'inspect', 'columns', 'preview', 'import'):
            response = FlaskClient(app.app).post('/api/spreadsheet/' + path, json={}, headers=HEADERS)
            self.assertEqual(response.status_code, 403, (path, response.json))
        self.assertEqual(self.status()['connection'], original['connection'])
        self.assertEqual(self.status()['source'], original['source'])
        self.assertEqual(self.post('source', {'link': LINK}).status_code, 200)

    def test_reset_cannot_disable_connection_without_admin(self):
        self.connect()
        response = FlaskClient(app.app).post('/api/dashboard/reset', json={
            'confirmed_reset': True, 'confirmation': 'RESET DASHBOARD'})
        self.assertEqual(response.status_code, 403)
        self.assertIsNotNone(self.status()['connection'])
        self.assertEqual(len(self.status()['leads']), 1)

    def test_unconfigured_admin_cannot_change_source(self):
        with patch.object(app, 'admin_password_hash', return_value=''):
            self.assertEqual(FlaskClient(app.app).post('/api/spreadsheet/source', json={'link': LINK}, headers=HEADERS).status_code, 403)

    def test_wrong_password_cannot_replace_saved_source(self):
        self.assertEqual(self.post('source', {'link': LINK}).status_code, 200)
        original = self.status()['source']
        response = FlaskClient(app.app).post('/api/spreadsheet/source', json={
            'link': LINK.replace('private_sheet', 'different_sheet'), 'admin_password': 'wrong'}, headers=HEADERS)
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.status()['source'], original)

    def test_links_only_allow_known_google_hosts_and_file_ids(self):
        self.assertEqual(parse_link(LINK), 'private_sheet_123456')
        self.assertEqual(parse_link('https://drive.google.com/file/d/private_sheet_123456/view'), 'private_sheet_123456')
        self.assertEqual(parse_link('https://drive.google.com/open?id=private_sheet_123456'), 'private_sheet_123456')
        self.assertEqual(parse_link('https://docs.google.com/spreadsheets/u/0/d/private_sheet_123456/edit'), 'private_sheet_123456')
        for value in ('http://localhost/data', 'https://docs.google.com.evil.test/spreadsheets/d/private_sheet_123456',
                      'https://evil@docs.google.com/spreadsheets/d/private_sheet_123456',
                      'https://docs.google.com:bad/spreadsheets/d/private_sheet_123456', 'https://[bad', 'not a link', None):
            with self.subTest(value=value), self.assertRaises(SheetError):
                parse_link(value)

    def test_private_access_and_excel_file_rejection(self):
        provider = GoogleSheets()
        session = unittest.mock.MagicMock()
        session.__enter__.return_value = session
        with patch.object(provider, 'session', return_value=session), patch.object(provider, 'get', return_value={'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}):
            with self.assertRaisesRegex(SheetError, 'Save as Google Sheets'):
                provider.inspect('private_sheet_123456')
        response = unittest.mock.Mock(status_code=403)
        session.get.return_value = response
        with self.assertRaisesRegex(SheetError, 'Viewer'):
            provider.get(session, 'https://sheets.googleapis.com/test', {})
        self.provider.failure = SheetError('No permission. Share as Viewer.', 403)
        result = self.post('inspect', {'link': LINK})
        self.assertEqual(result.status_code, 403)
        self.assertIsNone(self.status()['connection'])

    def test_preview_is_read_only_and_import_uses_exact_reviewed_snapshot(self):
        preview = self.post('preview', connection_payload()).json
        self.assertEqual(preview['report']['counts']['leads']['added'], 1)
        self.assertEqual(self.status()['leads'], [])
        self.provider.rows[0][1][1] = 'Changed after preview'
        imported = self.post('import', {'token': preview['token']}).json
        self.assertEqual(imported['leads'][0]['name'], 'Hospital A')
        self.assertEqual(imported['leads'][0]['id'], '0001')
        self.assertEqual(imported['leads'][0]['mobile'], '0012345678')
        self.assertEqual(imported['queries'][0]['lead_id'], '0001')
        self.assertEqual(self.post('import', {'token': preview['token']}).status_code, 409)

    def test_unchanged_sync_does_not_download_the_full_crm_snapshot(self):
        self.connect()
        before = self.client.get('/api/workspace').json['workspace']['revision']
        # Workbook/backups live in this snapshot; unchanged imports must not read it.
        with patch('shared_workspace.read', side_effect=AssertionError('Full snapshot read on unchanged sync')):
            response = self.post('sync')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['connection']['report']['counts']['leads']['unchanged'], 1)
        self.assertEqual(self.client.get('/api/workspace').json['workspace']['revision'], before)

    def test_deletion_preview_is_read_only_and_query_deletion_keeps_lead(self):
        self.connect()
        self.provider.rows[1] = [self.provider.rows[1][0]]
        preview = self.post('preview', connection_payload()).json
        self.assertEqual(preview['report']['counts']['queries']['removed'], 1)
        self.assertEqual(len(self.status()['queries']), 1)
        result = self.post('import', {'token': preview['token']}).json
        self.assertEqual(len(result['leads']), 1)
        self.assertEqual(result['queries'], [])
        self.assertEqual(len(self.client.get('/api/workspace').json['workspace']['values']['hk_v2_leads']), 1)

    def test_delete_lead_removes_linked_query_but_keeps_manual_records(self):
        self.connect()
        workspace = self.client.get('/api/workspace').json['workspace']
        leads = workspace['values']['hk_v2_leads'] + [{'id': 'MANUAL', 'name': 'Manual hospital'}]
        self.client.post('/api/workspace', json={'revision': workspace['revision'], 'values': {
            'hk_v2_leads': leads,
            'hk_v2_clients': [{'id': '0001', 'name': 'Unrelated manual client'}],
            'hk_v2_potential': [{'id': 'COPY', 'name': 'Linked hospital', 'sourceWorkspace': 'leads', 'sourceRecordId': '0001'}]
        }}, headers=HEADERS)
        self.provider.rows[0] = [self.provider.rows[0][0]]
        result = self.post('sync')
        self.assertEqual(result.status_code, 200, result.json)
        self.assertEqual(result.json['queries'], [])
        values = self.client.get('/api/workspace').json['workspace']['values']
        self.assertEqual([row['id'] for row in values['hk_v2_leads']], ['MANUAL'])
        self.assertEqual(values['hk_v2_clients'][0]['name'], 'Unrelated manual client')
        self.assertEqual(values['hk_v2_potential'], [])
        self.assertEqual(values['hk_v2_segments']['potential'], [])
        self.assertEqual(result.json['connection']['report']['counts']['queries']['removed'], 1)
        repeat = self.post('sync').json
        self.assertEqual(repeat['connection']['report']['counts']['leads']['removed'], 0)

    def test_invalid_rows_and_failed_reads_do_not_cause_deletions(self):
        self.connect()
        self.provider.rows[0][1][0] = ''
        result = self.post('sync').json
        self.assertEqual(len(result['leads']), 1)
        self.assertEqual(result['connection']['report']['counts']['leads']['removed'], 0)
        self.assertTrue(result['connection']['report']['warnings'])
        self.provider.failure = SheetError('Unavailable', 502)
        self.assertEqual(self.post('sync').status_code, 502)
        self.assertEqual(len(self.status()['leads']), 1)

    def test_source_switch_does_not_delete_other_source_records(self):
        self.connect()
        self.provider.rows[0][1][0] = 'OTHER'
        self.provider.rows[1] = [self.provider.rows[1][0]]
        inspect = self.provider.inspect
        def inspect_source(file_id):
            info = inspect(file_id)
            info['url'] = 'https://docs.google.com/spreadsheets/d/' + file_id + '/edit'
            return info
        self.provider.inspect = inspect_source
        payload = connection_payload();payload['link'] = LINK.replace('private_sheet_123456', 'other_sheet_123456')
        self.connect(payload)
        self.assertEqual(len(self.status()['leads']), 2)
        self.provider.rows[0] = [self.provider.rows[0][0]]
        result = self.post('sync').json
        self.assertEqual([row['id'] for row in result['leads']], ['0001'])
        self.assertEqual(len(result['queries']), 1)

    def test_deletion_failure_rolls_back_source_and_dashboard_together(self):
        before = self.connect()
        workspace = self.client.get('/api/workspace').json['workspace']
        self.provider.rows = [[self.provider.rows[0][0]], [self.provider.rows[1][0]]]
        with patch('shared_workspace.remove_sheet_leads', side_effect=RuntimeError('simulated failure')):
            self.assertEqual(self.post('sync').status_code, 502)
        after = self.status()
        self.assertEqual(after['leads'], before['leads'])
        self.assertEqual(after['queries'], before['queries'])
        self.assertEqual(self.client.get('/api/workspace').json['workspace'], workspace)

    def test_legacy_connection_reconciles_previously_imported_rows(self):
        self.connect()
        with app.db_conn() as conn:
            import json
            row = conn.execute('SELECT settings FROM sheet_connection WHERE id=1').fetchone()
            settings = json.loads(row['settings']);settings.pop('synced_ids')
            conn.execute('UPDATE sheet_connection SET settings=? WHERE id=1', (json.dumps(settings),))
        self.provider.rows = [[self.provider.rows[0][0]], [self.provider.rows[1][0]]]
        result = self.post('sync').json
        self.assertEqual(result['leads'], [])
        self.assertEqual(result['queries'], [])

    def test_add_edit_repeat_and_missing_rows_remove_synced_records(self):
        self.connect()
        self.provider.rows[0].append(['0002', 'Hospital B', '00009', 'New Lead', '2026-10-02', ''])
        self.provider.rows[0][1][1] = 'Hospital A updated'
        first = self.post('sync').json
        self.assertEqual(first['connection']['report']['counts']['leads'], {'added': 1, 'updated': 1, 'removed': 0, 'unchanged': 0, 'rejected': 0})
        second = self.post('sync').json
        self.assertEqual(len(second['leads']), 2)
        self.assertEqual(second['connection']['report']['counts']['leads']['unchanged'], 2)
        self.assertEqual(second['connection']['report']['counts']['queries']['unchanged'], 1)
        self.provider.rows = [[self.provider.rows[0][0]], [self.provider.rows[1][0]]]
        removed = self.post('sync').json
        self.assertEqual(len(removed['leads']), 0)
        self.assertEqual(len(removed['queries']), 0)
        self.assertEqual(removed['connection']['report']['counts']['leads']['removed'], 2)
        self.assertEqual(self.client.get('/api/workspace').json['workspace']['values']['hk_v2_leads'], [])

    def test_row_errors_reject_duplicates_missing_ids_bad_dates_and_orphans(self):
        self.provider.rows[0] += [
            ['0001', 'Duplicate', '', '', '', ''],
            ['', 'Missing ID'], ['0003', 'Bad date', '', '', '22/10/2026'],
            ['0004', 'Unknown status', '', 'made-up'], ['bad\'id', 'Unsafe ID'],
            ['0005', 'Valid lead', '', 'New Lead', '', ''],
        ]
        self.provider.rows[1] += [
            ['0008', '0005', 'Valid query'], ['0009', 'nonexistent', 'Orphan'],
            ['0010', '0005', 'Dup query'], ['0010', '0005', 'Dup query 2'],
        ]
        result = self.connect()
        counts = result['connection']['report']['counts']
        self.assertEqual(counts['leads']['rejected'], 6)
        self.assertEqual(counts['queries']['rejected'], 4)
        self.assertEqual([row['id'] for row in result['leads']], ['0005'])
        self.assertEqual([row['id'] for row in result['queries']], ['0008'])

    def test_unmapped_fields_survive_and_mapped_blanks_clear(self):
        self.connect()
        payload = connection_payload()
        del payload['mappings']['leads']['mobile']
        self.provider.rows[0][1][2] = 'different'
        self.assertEqual(self.connect(payload)['leads'][0]['mobile'], '0012345678')
        self.provider.rows[0][1][2] = ''
        self.assertEqual(self.connect()['leads'][0]['mobile'], '')

    def test_structural_and_connection_failures_preserve_last_good_data(self):
        initial = self.connect()
        self.provider.rows[0][0][0] = 'Renamed ID header'
        response = self.post('sync')
        self.assertEqual(response.status_code, 400)
        state = self.status()
        self.assertEqual(state['leads'], initial['leads'])
        self.assertEqual(state['connection']['last_success'], initial['connection']['last_success'])
        self.assertIn('Missing or duplicate column', state['connection']['last_error'])
        self.provider.failure = SheetError('Google connection failed.', 502)
        self.assertEqual(self.post('sync').status_code, 502)
        self.assertEqual(self.status()['queries'], initial['queries'])

    def test_bad_mappings_and_same_tabs_are_clear_errors(self):
        for modification in ('missing', 'duplicate', 'same-tab', 'invalid-field', 'invalid-value'):
            payload = connection_payload()
            if modification == 'missing': del payload['mappings']['leads']['id']
            if modification == 'duplicate': payload['mappings']['leads']['mobile'] = 'Name'
            if modification == 'same-tab': payload['tabs']['queries'] = 'Sales'
            if modification == 'invalid-field': payload['mappings']['leads']['password'] = 'Name'
            if modification == 'invalid-value': payload['mappings']['leads']['id'] = []
            with self.subTest(modification=modification):
                self.assertEqual(self.post('preview', payload).status_code, 400)

    def test_atomic_database_failure_rolls_back_both_entity_types(self):
        self.connect()
        before = self.status()
        self.provider.rows[0][1][1] = 'Would change'
        self.provider.rows[1][1][2] = 'Fail query'
        with app.db_conn() as conn:
            conn.execute("CREATE TRIGGER reject_query BEFORE UPDATE ON sheet_queries BEGIN SELECT RAISE(ABORT, 'test failure'); END")
        self.assertEqual(self.post('sync').status_code, 502)
        after = self.status()
        self.assertEqual(before['leads'], after['leads'])
        self.assertEqual(before['queries'], after['queries'])
        self.assertEqual(before['connection']['revision'], after['connection']['revision'])

    def test_preview_revision_guard_and_expiry(self):
        self.connect()
        token = self.post('preview', connection_payload()).json['token']
        self.post('sync')
        self.assertEqual(self.post('import', {'token': token}).status_code, 409)
        token = self.post('preview', connection_payload()).json['token']
        with app.db_conn() as conn:
            conn.execute('UPDATE sheet_previews SET expires=0 WHERE token=?', (token,))
        self.assertEqual(self.post('import', {'token': token}).status_code, 409)

    def test_import_failure_returns_clear_error_and_saves_no_partial_records(self):
        token = self.post('preview', connection_payload()).json['token']
        with app.db_conn() as conn:
            conn.execute("CREATE TRIGGER reject_new_query BEFORE INSERT ON sheet_queries BEGIN SELECT RAISE(ABORT, 'test failure'); END")
        result = self.post('import', {'token': token})
        self.assertEqual(result.status_code, 503)
        self.assertIn('No partial import', result.json['error'])
        self.assertEqual(self.status()['leads'], [])
        self.assertIsNone(self.status()['connection'])

    def test_connection_persists_and_auto_refresh_runs_only_when_due(self):
        payload = connection_payload(); payload['auto_minutes'] = 5
        self.connect(payload)
        app.init_db()
        self.assertEqual(self.status()['connection']['name'], 'Private Sales')
        self.provider.rows[0][1][1] = 'Automatic update'
        app.sheet_sync.auto_sync()
        self.assertEqual(self.status()['leads'][0]['name'], 'Hospital A')
        with app.db_conn() as conn:
            conn.execute('UPDATE sheet_connection SET last_attempt_epoch=0')
        app.sheet_sync.auto_sync()
        self.assertEqual(self.status()['leads'][0]['name'], 'Automatic update')

    def test_export_real_xlsx_text_dates_headers_and_sensitive_field_exclusion(self):
        payload = {'leads': [{'id': '0001', 'name': '=HYPERLINK("bad")', 'mobile': '00123',
                              'created_date': '2026-10-01', 'next': '2026-11-01',
                              'password': 'secret', 'access_token': 'secret'}],
                   'queries': [{'id': '0009', 'lead_id': '0001', 'subject': 'Monitor', 'updated_date': '2026-10-02T10:30:00+05:30'}]}
        response = self.post('export', payload)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data.startswith(b'PK'))
        workbook = load_workbook(BytesIO(response.data))
        self.assertEqual(workbook.sheetnames, ['Leads', 'Queries'])
        sheet = workbook['Leads']; query = workbook['Queries']
        self.assertEqual(sheet['A2'].value, '0001')
        self.assertEqual(sheet['A2'].number_format, '@')
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(sheet['E2'].value, '00123')
        self.assertEqual(sheet['I2'].value.date().isoformat(), '2026-10-01')
        self.assertEqual(query['B2'].value, '0001')
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertNotIn('password', [cell.value for cell in sheet[1]])
        self.assertNotIn('secret', str(list(sheet.values)))
        self.assertEqual(self.post('export', {'leads': [], 'queries': []}).status_code, 200)

    def test_cross_origin_writes_and_reset_keep_connection_but_pause_auto(self):
        self.connect()
        response = self.client.post('/api/spreadsheet/sync', json={}, headers={**HEADERS, 'Origin': 'https://evil.test'})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.post('/api/spreadsheet/sync', json={}).status_code, 403)
        response = self.client.post('/api/dashboard/reset', json={'confirmed_reset': True, 'confirmation': 'RESET DASHBOARD', 'admin_password': 'sheet-test-admin'})
        self.assertEqual(response.status_code, 200)
        status = self.status()
        self.assertEqual(status['leads'], [])
        self.assertEqual(status['queries'], [])
        self.assertEqual(status['connection']['auto_minutes'], 0)
        self.assertIsNone(status['connection']['last_success'])


if __name__ == '__main__':
    unittest.main()
