"""Chrome/Edge workflow checks with a temporary DB and a synthetic private sheet.

No Google account, live CRM data, email or WhatsApp operations are used.
"""
import json
import os
import sys
import tempfile
from pathlib import Path
from threading import Thread
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from waitress import create_server
from test_spreadsheet_sync import FakeSheets, SheetError


def check():
    with tempfile.TemporaryDirectory() as temp, patch.object(app, 'DB_PATH', Path(temp) / 'test.db'), patch.object(app, 'DATA_DIR', Path(temp)):
        app.init_db()
        if app.scheduler.running:
            app.scheduler.pause()
        server = create_server(app.app, host='127.0.0.1', port=0)
        Thread(target=server.run, daemon=True).start()
        reports = []
        try:
            with sync_playwright() as playwright:
                for channel in os.getenv('BROWSER_CHANNELS', 'chrome,msedge').split(','):
                    # Each browser's fixture and server records start independently.
                    with app.db_conn() as conn:
                        for table in ('sheet_queries', 'sheet_leads', 'sheet_previews', 'sheet_connection', 'crm_workspace'):
                            conn.execute('DELETE FROM ' + table)
                    app.init_db()
                    provider = FakeSheets()
                    with patch.object(app.sheet_sync, 'provider', provider):
                        browser = playwright.chromium.launch(channel=None if channel == 'chromium' else channel, headless=True)
                        context = browser.new_context(accept_downloads=True, viewport={'width': 1440, 'height': 1000})
                        page = context.new_page(); errors = []
                        page.on('pageerror', lambda error: errors.append(str(error)))
                        page.goto(f'http://127.0.0.1:{server.effective_port}/')
                        page.wait_for_function('dashboardReady && !document.body.inert && sheetState !== null')
                        base = page.evaluate('({leads:sectionRecords.leads.length,clients:JSON.stringify(clients),total:workspaceRecords().length})')
                        page.evaluate('openSpreadsheetConnection()')
                        page.locator('#sheetLinkInput').fill('https://docs.google.com/spreadsheets/d/private_sheet_123456/edit')
                        page.get_by_role('button', name='Check link and access', exact=True).click()
                        page.wait_for_function('!sheetBusy && sheetInspection !== null')
                        page.locator('#sheetLeadsTab').select_option('Sales')
                        page.locator('#sheetQueriesTab').select_option('Enquiries')
                        page.get_by_role('button', name='Load columns', exact=True).click()
                        page.wait_for_function("!sheetBusy && !document.getElementById('sheetMappingStep').classList.contains('hidden')")
                        page.locator('#sheetMap-queries-updated_date').select_option('Updated')
                        page.get_by_role('button', name='Preview and validate', exact=True).click()
                        page.wait_for_function('!sheetBusy && sheetPreviewToken !== null')
                        assert 'Hospital A' in page.locator('#sheetPreviewRows').inner_text()
                        assert page.evaluate('sectionRecords.leads.length') == base['leads'], 'Preview must not import'
                        page.locator('#sheetImportButton').click()
                        page.wait_for_function("!sheetBusy && sectionRecords.leads.some(row=>row.id==='0001')")
                        assert page.locator('#sheetLeadRows').inner_text().count('0012345678') == 1
                        assert '0007' in page.locator('#sheetQueryRows').inner_text()
                        assert page.locator('#leadTotal').inner_text() == str(base['leads'] + 1)
                        assert page.locator('#kpiClients').inner_text() == str(base['total'] + 1)
                        assert page.evaluate('JSON.stringify(clients)') == base['clients']
                        assert page.locator('#sheetOpenLink').get_attribute('href').startswith('https://docs.google.com/spreadsheets/')

                        provider.rows[0][1][1] = 'Hospital A updated'
                        provider.rows[0].append(['0002', 'Hospital B', '00009', 'New Lead', '2026-10-02', ''])
                        provider.rows[1][1][3] = 'Revised requirements'
                        page.locator('#sheetSyncButton').click()
                        page.wait_for_function("!sheetBusy && sectionRecords.leads.some(row=>row.name==='Hospital A updated')")
                        page.locator('#sheetSyncButton').click()
                        page.wait_for_function('!sheetBusy')
                        assert page.evaluate("sectionRecords.leads.filter(row=>row.id==='0002').length") == 1
                        assert '2 unchanged' in page.locator('#sheetSyncReport').inner_text()
                        assert 'Revised requirements' in page.locator('#sheetQueryRows').inner_text()
                        assert page.locator('#leadTotal').inner_text() == str(base['leads'] + 2)
                        assert page.locator('#kpiClients').inner_text() == str(base['total'] + 2)

                        def download(scope, expected_leads, expected_queries, name):
                            page.evaluate('openLeadQueryExport()')
                            page.locator('#sheetExportScope').select_option(scope)
                            with page.expect_download() as event:
                                page.locator('#sheetDownloadButton').click()
                            target = Path(temp) / f'{channel}-{name}.xlsx'
                            event.value.save_as(target)
                            workbook = load_workbook(target)
                            assert workbook.sheetnames == ['Leads', 'Queries']
                            assert workbook['Leads'].max_row - 1 == expected_leads
                            assert workbook['Queries'].max_row - 1 == expected_queries
                            return workbook

                        page.locator('#sheetLeadSearch').fill('Hospital A updated')
                        page.locator('#sheetQuerySearch').fill('Revised requirements')
                        workbook = download('filtered', 1, 1, 'filtered')
                        assert workbook['Leads']['A2'].value == '0001'
                        assert workbook['Leads']['E2'].value == '0012345678'
                        query_total = page.evaluate('allQueryRecords().length')
                        download('all', base['leads'] + 2, query_total, 'all')

                        page.evaluate("navigate('dashboard')")
                        page.locator('#search').fill('Hospital A updated')
                        download('filtered', 1, 1, 'dashboard-filtered')

                        page.evaluate("navigate('leads')")
                        before_failure = page.evaluate('JSON.stringify({leads:sectionRecords.leads,queries:allQueryRecords()})')
                        provider.failure = SheetError('Permission removed. Share as Viewer.', 403)
                        page.locator('#sheetSyncButton').click()
                        page.wait_for_function('!sheetBusy')
                        assert 'Permission removed' in page.locator('#sheetConnectionError').inner_text()
                        assert before_failure == page.evaluate('JSON.stringify({leads:sectionRecords.leads,queries:allQueryRecords()})')
                        provider.failure = None
                        page.reload()
                        page.wait_for_function('dashboardReady && !document.body.inert && sheetState?.connection !== null')
                        assert page.evaluate('sectionRecords.leads.length') == base['leads'] + 2
                        assert page.evaluate('sheetState.queries.length') == 1

                        # A second browser profile receives the server-saved records, too.
                        other_context = browser.new_context()
                        other_page = other_context.new_page(); other_page.goto(page.url)
                        other_page.wait_for_function('dashboardReady && !document.body.inert && sheetState !== null')
                        assert other_page.evaluate("sectionRecords.leads.some(row=>row.id==='0001'&&row.name==='Hospital A updated')")
                        other_context.close()

                        page.set_viewport_size({'width': 390, 'height': 844})
                        page.evaluate('openSpreadsheetConnection()')
                        assert page.locator('#sheetLinkInput').is_visible()
                        assert page.locator('#sheetSyncButton').is_visible()
                        assert not errors, errors
                        reports.append({'browser': channel, 'add_update_idempotency': True, 'all_and_filtered_xlsx': True,
                                        'retained_on_failure': True, 'reload_and_profile_persistence': True,
                                        'dashboard_counts': True, 'mobile_controls': True, 'javascript_errors': errors})
                        context.close(); browser.close()
        finally:
            server.close()
        print(json.dumps(reports, indent=2))


if __name__ == '__main__':
    check()
