"""Shared Drive bookmarks and authorized, bounded workbook downloads for import."""
import re
import secrets
import sqlite3
from io import BytesIO
from urllib.parse import parse_qs, urlencode, urlparse

from flask import Blueprint, jsonify, request, send_file
from request_security import crm_request_origin_allowed
from spreadsheet_sync import GoogleSheets, SheetError, parse_link

TARGETS = {'clients': 'Hospitals & Clients', 'rghs': 'RGHS Data', 'served': 'Already Served',
           'doctors': 'Doctors', 'potential': 'Potential Leads', 'leads': 'Leads & Queries', 'folder': 'CRM source folder'}
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MAX_BYTES = 16 * 1024 * 1024


def parse_drive_reference(link):
    if not isinstance(link, str) or len(link) > 2048:
        raise SheetError('Enter a Google Drive folder, workbook or Sheets link.')
    try:
        url = urlparse(link.strip())
        valid = url.scheme == 'https' and url.hostname == 'drive.google.com' and not url.username and not url.password and url.port in (None,443)
        folder = re.fullmatch(r'/drive/(?:u/\d+/)?folders/([A-Za-z0-9_-]{10,200})/?',url.path) if valid else None
    except ValueError:
        raise SheetError('Invalid Google Drive link.') from None
    return (folder[1],True) if folder else (parse_link(link),False)


def init_schema(conn):
    conn.execute('''CREATE TABLE IF NOT EXISTS shared_drive_links (
        id TEXT PRIMARY KEY, file_id TEXT NOT NULL, link TEXT NOT NULL,
        label TEXT NOT NULL, target TEXT NOT NULL, created_at TEXT NOT NULL,
        UNIQUE(file_id,target))''')
    conn.executescript('''CREATE TABLE IF NOT EXISTS drive_workbook_copies (
        file_id TEXT PRIMARY KEY, filename TEXT NOT NULL, mime TEXT NOT NULL,
        content BLOB NOT NULL, saved_at TEXT NOT NULL, size_bytes INTEGER NOT NULL
    );''')


class DriveWorkbook(GoogleSheets):
    def session(self):
        try:
            import google.auth
            from google.auth.transport.requests import AuthorizedSession
            credentials, _ = google.auth.default(scopes=['https://www.googleapis.com/auth/drive.readonly'])
            return AuthorizedSession(credentials, refresh_timeout=15)
        except Exception:
            raise SheetError('Google access is not configured on this CRM server. Configure private server credentials and share the Drive file with that service account as Viewer. You can still save/open the link or download the workbook from Drive and use Import Excel.', 503) from None

    def download(self, file_id):
        with self.session() as session:
            metadata = self.get(session, 'https://www.googleapis.com/drive/v3/files/' + file_id,
                {'fields': 'id,name,mimeType,capabilities(canDownload)', 'supportsAllDrives': 'true'})
            if metadata.get('capabilities', {}).get('canDownload') is False:
                raise SheetError('The file owner has disabled downloading. Ask the owner to allow workbook download.', 403)
            mime = metadata.get('mimeType')
            native = mime == 'application/vnd.google-apps.spreadsheet'
            if not native and mime not in (XLSX, 'application/vnd.ms-excel'):
                raise SheetError('Unsupported Drive file. Choose a native Google Sheet or an Excel .xlsx/.xls workbook. Folders, shortcuts, PDFs and Google Forms are not workbook imports.')
            suffix = '.xls' if mime == 'application/vnd.ms-excel' else '.xlsx'
            name = re.sub(r'[\x00-\x1f\x7f/\\]', '_', metadata.get('name') or 'Drive workbook')[:180]
            if not name.lower().endswith(suffix): name += suffix
            endpoint = 'https://www.googleapis.com/drive/v3/files/' + file_id
            params = {'alt': 'media', 'supportsAllDrives': 'true'}
            if native: endpoint += '/export'; params = {'mimeType': XLSX}
            try:
                with session.get(endpoint, params=params, stream=True, timeout=(10, 45), allow_redirects=False) as response:
                    if response.status_code in (401, 403, 404):
                        raise SheetError('Cannot download this workbook. Check Viewer sharing, Google Drive API access, and the drive.readonly credential scope. Native Sheets exports must also fit Google’s 10 MB export limit.', 403)
                    if response.status_code == 429:
                        raise SheetError('Google request limit reached. Wait a minute and retry.', 503)
                    if response.status_code != 200:
                        raise SheetError('Google could not download the workbook. Existing CRM data was kept.', 502)
                    content = BytesIO()
                    for chunk in response.iter_content(64 * 1024):
                        if content.tell() + len(chunk) > MAX_BYTES:
                            raise SheetError('The workbook exceeds the 16 MB import limit. Split the source workbook.', 413)
                        content.write(chunk)
                    if content.tell() == 0: raise SheetError('Google returned an empty workbook.', 502)
                    content.seek(0)
                    return content, name, suffix
            except SheetError:
                raise
            except Exception:
                raise SheetError('Drive download failed or timed out. Existing saved records are unchanged.', 502) from None


def register(app, db_conn, now_iso, storage_info=None):
    routes = Blueprint('drive_import', __name__, url_prefix='/api/drive-import')
    provider = DriveWorkbook()

    @routes.before_request
    def guard():
        if request.method == 'POST':
            copy_upload = request.endpoint == 'drive_import.save_copy'
            if not crm_request_origin_allowed(request) or request.headers.get('X-CRM-Request') != '1' or (not copy_upload and not request.is_json) or (copy_upload and request.mimetype != 'multipart/form-data'):
                raise SheetError('Use this CRM dashboard to manage Drive links.', 403)
            if request.content_length and request.content_length > (MAX_BYTES + 8192 if copy_upload else 8192):
                raise SheetError('Drive link request is too large.', 413)

    @routes.after_request
    def no_cache(response):
        response.headers['Cache-Control'] = 'no-store, private'
        return response

    @routes.errorhandler(SheetError)
    def error(exc): return jsonify(ok=False, error=str(exc)), exc.status

    @routes.errorhandler(sqlite3.Error)
    def database_error(exc): return jsonify(ok=False, error='Shared Drive links could not be saved. Existing links and records were kept.'), 503

    def body():
        data = request.get_json(silent=True)
        if not isinstance(data, dict): raise SheetError('Send a Drive link request object.')
        return data

    def state():
        with db_conn() as conn:
            rows = [dict(row) for row in conn.execute('''SELECT links.id,links.link,links.label,links.target,links.created_at,
                copies.saved_at AS copy_saved_at,copies.filename AS copy_filename,copies.size_bytes AS copy_size
                FROM shared_drive_links links LEFT JOIN drive_workbook_copies copies ON copies.file_id=links.file_id
                ORDER BY links.created_at,links.id''')]
        return jsonify(ok=True, links=rows, storage=storage_info() if storage_info else {'engine': 'sqlite', 'durable': False})

    @routes.get('')
    def status(): return state()

    @routes.post('/save')
    def save():
        data = body(); file_id, folder = parse_drive_reference(data.get('link'))
        target, label = data.get('target', 'clients'), data.get('label', '')
        if not isinstance(target, str) or target not in TARGETS: raise SheetError('Choose a valid CRM import section.')
        if not isinstance(label, str) or not label.strip() or len(label) > 120:
            raise SheetError('Enter a link label of 1–120 characters.')
        if folder: target = 'folder'
        elif target == 'folder': raise SheetError('Choose a Drive folder for the folder destination.')
        link = 'https://drive.google.com/drive/folders/' + file_id if folder else 'https://docs.google.com/spreadsheets/d/' + file_id + '/edit' if '/spreadsheets/' in data['link'] else 'https://drive.google.com/file/d/' + file_id + '/view'
        resource_key = parse_qs(urlparse(data['link']).query).get('resourcekey', [''])[0]
        if resource_key and re.fullmatch(r'[A-Za-z0-9_-]{1,200}', resource_key):
            link += '?' + urlencode({'resourcekey': resource_key})
        with db_conn() as conn:
            if conn.execute('SELECT count(*) FROM shared_drive_links').fetchone()[0] >= 100 and not conn.execute('SELECT 1 FROM shared_drive_links WHERE file_id=? AND target=?', (file_id, target)).fetchone():
                raise SheetError('At most 100 Drive links can be saved. Remove an unused link first.')
            conn.execute('INSERT INTO shared_drive_links VALUES (?,?,?,?,?,?) ON CONFLICT(file_id,target) DO UPDATE SET label=excluded.label,link=excluded.link',
                         (secrets.token_urlsafe(18), file_id, link, label.strip(), target, now_iso()))
        return state()

    @routes.post('/remove')
    def remove():
        identifier = body().get('id')
        if not isinstance(identifier, str): raise SheetError('Choose a saved link.')
        with db_conn() as conn: conn.execute('DELETE FROM shared_drive_links WHERE id=?', (identifier,))
        return state()

    @routes.post('/download')
    def download():
        data = body(); identifier = data.get('id')
        if not isinstance(identifier, str): raise SheetError('Choose a saved link.')
        with db_conn() as conn: row = conn.execute('SELECT file_id,target FROM shared_drive_links WHERE id=?', (identifier,)).fetchone()
        if not row: raise SheetError('This shared link was removed. Refresh the list.', 404)
        if row['target'] == 'folder': raise SheetError('Open the folder and choose an Excel or Sheets file to import.')
        if data.get('saved_copy') is True:
            with db_conn() as conn: copy = conn.execute('SELECT filename,mime,content FROM drive_workbook_copies WHERE file_id=?', (row['file_id'],)).fetchone()
            if not copy: raise SheetError('No saved workbook copy is available. Attach an Excel copy or configure Google access.', 404)
            return send_file(BytesIO(bytes(copy['content'])), mimetype=copy['mime'], as_attachment=True, download_name=copy['filename'])
        content, name, suffix = provider.download(row['file_id'])
        return send_file(content, mimetype=XLSX if suffix == '.xlsx' else 'application/vnd.ms-excel', as_attachment=True, download_name=name)

    @routes.post('/save-copy')
    def save_copy():
        identifier = request.form.get('id')
        with db_conn() as conn: row = conn.execute('SELECT file_id,target FROM shared_drive_links WHERE id=?', (identifier,)).fetchone()
        if not row: raise SheetError('Choose a saved Drive link before attaching its Excel copy.', 404)
        if row['target'] == 'folder': raise SheetError('Attach copies to workbook links, not folders.')
        file = request.files.get('file')
        if not file or not file.filename: raise SheetError('Choose the downloaded Excel workbook for this Drive link.')
        name = re.sub(r'[\x00-\x1f\x7f/\\]', '_', file.filename)[:180]
        suffix = '.xlsx' if name.lower().endswith('.xlsx') else '.xls' if name.lower().endswith('.xls') else None
        if not suffix: raise SheetError('Attach an .xlsx or .xls workbook copy.')
        content = file.stream.read(MAX_BYTES + 1)
        if len(content) > MAX_BYTES: raise SheetError('Workbook copies must be 16 MB or smaller.', 413)
        try:
            if suffix == '.xlsx':
                from openpyxl import load_workbook
                workbook = load_workbook(BytesIO(content), read_only=True, data_only=True); workbook.close()
            else:
                import xlrd
                workbook = xlrd.open_workbook(file_contents=content, on_demand=True); workbook.release_resources()
        except Exception:
            raise SheetError('The attached file is not a readable Excel workbook. The previous saved copy was kept.') from None
        with db_conn() as conn:
            conn.execute('INSERT INTO drive_workbook_copies VALUES (?,?,?,?,?,?) ON CONFLICT(file_id) DO UPDATE SET filename=excluded.filename,mime=excluded.mime,content=excluded.content,saved_at=excluded.saved_at,size_bytes=excluded.size_bytes',
                         (row['file_id'],name,XLSX if suffix=='.xlsx' else 'application/vnd.ms-excel',content,now_iso(),len(content)))
        return state()

    app.register_blueprint(routes)
    return provider
