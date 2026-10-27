"""One-way private Google Sheets import and safe Leads/Queries XLSX export.

Credentials stay in Google ADC/the server environment. SQLite holds only connection
settings, reviewed snapshots, sync reports and imported business records.
"""
import json
import os
import re
import secrets
import sqlite3
from collections import Counter
from datetime import datetime, timezone
from io import BytesIO
from threading import RLock
from urllib.parse import parse_qs, urlparse

from flask import Blueprint, jsonify, request, send_file
import shared_workspace
import form_responses
from public_form_sheet import PublicFormSheet
from request_security import crm_request_origin_allowed

MAX_ROWS = 20000
MAX_COLUMNS = 200
MAX_BYTES = 12 * 1024 * 1024
ID_PATTERN = re.compile(r"[A-Za-z0-9_.:-]{1,128}\Z")
LEAD_STATUSES = ('New Lead', 'Contacted', 'Query Received', 'Quotation Shared',
                 'Follow-up Required', 'Not Interested', 'Negotiation',
                 'PO / Order Confirmed', 'Demand Fulfilled', 'Order Finalized',
                 'Delivered', 'Feedback Pending', 'Completed / Served',
                 'Lost / Cancelled', 'On Hold')
FIELDS = {
    'leads': [('id', 'Lead ID'), ('name', 'Name'), ('company', 'Company'),
              ('email', 'Email'), ('mobile', 'Phone'), ('source', 'Source'),
              ('status', 'Status'), ('executive', 'Assigned Owner'),
              ('created_date', 'Created Date'), ('next', 'Follow-up Date'),
              ('contact', 'Contact Person'), ('type', 'Client Type'),
              ('city', 'City'), ('district', 'District'), ('state', 'State'),
              ('requirement', 'Requirement'), ('product', 'Product'),
              ('quote', 'Quotation Value'), ('order', 'Order Value'),
              ('last', 'Last Contact Date'), ('priority', 'Priority'), ('remark', 'Remark')],
    'queries': [('id', 'Query ID'), ('lead_id', 'Related Lead ID'),
                ('subject', 'Subject'), ('details', 'Query Details'),
                ('status', 'Status'), ('executive', 'Assigned Owner'),
                ('created_date', 'Created Date'), ('updated_date', 'Last Updated Date')],
}
REQUIRED = {'leads': ('id', 'name'), 'queries': ('id', 'lead_id', 'subject')}
DATE_FIELDS = {'created_date', 'updated_date', 'next', 'last'}


class SheetError(ValueError):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def parse_link(link):
    if not isinstance(link, str) or len(link) > 2048:
        raise SheetError('Enter a Google Sheets or Google Drive spreadsheet link.')
    try:
        url = urlparse(link.strip())
    except ValueError:
        raise SheetError('Invalid Google spreadsheet link.') from None
    try:
        port = url.port
    except ValueError:
        raise SheetError('Invalid Google spreadsheet link.') from None
    if url.scheme != 'https' or url.username or url.password or port not in (None, 443):
        raise SheetError('Use an HTTPS Google Sheets or Google Drive link.')
    match = None
    if url.hostname == 'docs.google.com':
        match = re.fullmatch(r'/spreadsheets/(?:u/\d+/)?d/([A-Za-z0-9_-]{10,200})(?:/.*)?', url.path)
    elif url.hostname == 'drive.google.com':
        match = re.fullmatch(r'/file/d/([A-Za-z0-9_-]{10,200})(?:/.*)?', url.path)
        if not match and url.path in ('/open', '/file', '/file/'):
            file_id = parse_qs(url.query).get('id', [''])[0]
            if re.fullmatch(r'[A-Za-z0-9_-]{10,200}', file_id):
                return file_id
    if not match:
        raise SheetError('Invalid link. Paste the spreadsheet’s Sheets edit link or Drive file link.')
    return match.group(1)


class GoogleSheets:
    """Fixed Google API endpoints only; never fetch a user-provided URL."""
    def session(self):
        try:
            import google.auth
            from google.auth.transport.requests import AuthorizedSession
            credentials, _ = google.auth.default(scopes=[
                'https://www.googleapis.com/auth/spreadsheets.readonly',
                'https://www.googleapis.com/auth/drive.metadata.readonly',
            ])
            return AuthorizedSession(credentials, refresh_timeout=15)
        except Exception:
            raise SheetError('Google access is not configured. Ask the administrator to configure Google Application Default Credentials and share this sheet with the service account as Viewer.', 503) from None

    def get(self, session, url, params):
        try:
            response = session.get(url, params=params, timeout=(10, 45))
            if response.status_code in (401, 403, 404):
                raise SheetError('Cannot access this spreadsheet. Share it with the configured Google service account as Viewer, and enable the Google Sheets and Drive APIs. Check that the file still exists.', 403)
            if response.status_code == 429:
                raise SheetError('Google’s request limit was reached. Wait a minute and sync again.', 503)
            if response.status_code >= 400:
                raise SheetError('Google could not read the selected spreadsheet or tabs. Check access, file format and tab names.', 502)
            if len(response.content) > MAX_BYTES:
                raise SheetError('Spreadsheet response exceeds 12 MB. Split it into smaller tabs.')
            return response.json()
        except SheetError:
            raise
        except Exception:
            raise SheetError('Google connection failed or timed out. The last saved CRM data is retained.', 502) from None

    def inspect(self, file_id):
        with self.session() as session:
            file = self.get(session, 'https://www.googleapis.com/drive/v3/files/' + file_id,
                            {'fields': 'id,name,mimeType', 'supportsAllDrives': 'true'})
            if file.get('mimeType') != 'application/vnd.google-apps.spreadsheet':
                raise SheetError('Unsupported format: Excel files stored in Drive are not imported. Open the file in Google Sheets, choose File → Save as Google Sheets, and connect the new native Sheets link.')
            meta = self.get(session, 'https://sheets.googleapis.com/v4/spreadsheets/' + file_id,
                            {'fields': 'spreadsheetId,properties(title,locale,timeZone),sheets(properties(sheetId,title,sheetType,gridProperties))'})
            tabs = [dict(id=sheet['properties']['sheetId'], name=sheet['properties']['title'])
                    for sheet in meta.get('sheets', []) if sheet['properties'].get('sheetType', 'GRID') == 'GRID']
            if not tabs:
                raise SheetError('This spreadsheet has no supported grid worksheets.')
            return {'id': file_id, 'name': meta['properties']['title'], 'tabs': tabs,
                    'locale': meta['properties'].get('locale', 'en_US'), 'timezone': meta['properties'].get('timeZone', 'Asia/Kolkata'),
                    'url': 'https://docs.google.com/spreadsheets/d/' + file_id + '/edit'}

    def read(self, file_id, tabs, headers_only=False):
        end = '6' if headers_only else str(MAX_ROWS + 2)
        ranges = ["'" + title.replace("'", "''") + "'!A1:GS" + end for title in tabs]
        with self.session() as session:
            data = self.get(session, 'https://sheets.googleapis.com/v4/spreadsheets/' + file_id + '/values:batchGet',
                            [('ranges', value) for value in ranges] + [('valueRenderOption', 'FORMATTED_VALUE')])
        values = data.get('valueRanges', [])
        if len(values) != len(tabs):
            raise SheetError('Google returned an incomplete response. No records were imported.', 502)
        return [value.get('values', []) for value in values]


def init_schema(conn):
    conn.executescript('''
        CREATE TABLE IF NOT EXISTS sheet_connection (
            id INTEGER PRIMARY KEY CHECK(id=1), settings TEXT NOT NULL,
            revision INTEGER NOT NULL DEFAULT 0, last_success TEXT,
            last_attempt TEXT, last_error TEXT, last_result TEXT,
            last_attempt_epoch REAL NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS sheet_source_preferences (
            id INTEGER PRIMARY KEY CHECK(id=1), settings TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS sheet_leads (id TEXT PRIMARY KEY, data TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS sheet_queries (
            id TEXT PRIMARY KEY, lead_id TEXT NOT NULL REFERENCES sheet_leads(id), data TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS sheet_previews (
            token TEXT PRIMARY KEY, expires REAL NOT NULL, revision INTEGER NOT NULL, snapshot TEXT NOT NULL
        );
    ''')


def headers_for(rows):
    headers = [str(value).strip() for value in (rows[0] if rows else [])]
    if not headers or not any(headers):
        raise SheetError('The selected tab is empty. Put column headers in row 1.')
    if len(headers) > MAX_COLUMNS:
        raise SheetError('A selected tab exceeds 200 columns. Use a smaller source tab.')
    return headers


def normalize_date(value):
    if not value:
        return ''
    if not re.match(r'^\d{4}-\d{2}-\d{2}(?:$|T)', value):
        raise ValueError('Dates must use YYYY-MM-DD or an ISO timestamp; format the source cells as ISO dates.')
    return datetime.fromisoformat(value.replace('Z', '+00:00')).isoformat() if 'T' in value else datetime.strptime(value, '%Y-%m-%d').date().isoformat()


def validate_rows(kind, rows, mapping):
    headers = headers_for(rows)
    if len(rows) - 1 > MAX_ROWS:
        raise SheetError('A selected tab exceeds 20,000 rows. Split it before importing.')
    allowed = dict(FIELDS[kind])
    if not isinstance(mapping, dict) or any(key not in allowed or not isinstance(value, str) for key, value in mapping.items()):
        raise SheetError('Invalid column mapping.')
    for field in REQUIRED[kind]:
        if not mapping.get(field):
            raise SheetError('Map the required ' + allowed[field] + ' column.')
    used = [column for column in mapping.values() if column]
    if len(used) != len(set(used)):
        raise SheetError('Map each spreadsheet column to only one CRM field.')
    for column in used:
        if headers.count(column) != 1:
            raise SheetError('Missing or duplicate column: ' + str(column) + '. Update the column mapping.')
    mapped = {field: headers.index(column) for field, column in mapping.items() if column}
    nonempty = [(number, row) for number, row in enumerate(rows[1:], 2)
                if any(str(value).strip() for value in row)]
    ids = Counter(str(row[mapped['id']]).strip() for _, row in nonempty if len(row) > mapped['id'])
    accepted, errors = [], []
    for number, row in nonempty:
        data = {field: str(row[index]).strip() if index < len(row) else '' for field, index in mapped.items()}
        try:
            if not data.get('id'):
                raise ValueError('Missing stable ID.')
            if ids[data['id']] > 1:
                raise ValueError('Duplicate ID; every occurrence is rejected.')
            for field in ('id', 'lead_id'):
                if field in data and not ID_PATTERN.fullmatch(data[field]):
                    raise ValueError('IDs must contain 1–128 letters, digits, dots, colons, hyphens or underscores.')
            if any(not data.get(field) for field in REQUIRED[kind]):
                raise ValueError('A required field is empty.')
            if any(len(value) > 10000 for value in data.values()):
                raise ValueError('A cell exceeds 10,000 characters.')
            for field in DATE_FIELDS & data.keys():
                data[field] = normalize_date(data[field])
            if kind == 'leads':
                if data.get('status') and data['status'] not in LEAD_STATUSES:
                    raise ValueError('Unknown lead status. Use an existing CRM pipeline stage.')
                if data.get('email') and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', data['email']):
                    raise ValueError('Invalid email address.')
                for field in ('quote', 'order'):
                    if field in data:
                        value = float(data[field].replace(',', '') or '0')
                        if not 0 <= value <= 1e12:
                            raise ValueError('Quotation and order values must be finite, non-negative numbers.')
                        data[field] = value
            accepted.append({'row': number, 'data': data})
        except (ValueError, OverflowError) as exc:
            errors.append({'kind': kind, 'row': number, 'id': str(data.get('id', ''))[:128], 'reason': str(exc)})
    return accepted, errors


def json_text(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


class SheetSync:
    def __init__(self, db_conn, now_iso, provider=None):
        self.db_conn, self.now_iso = db_conn, now_iso
        self.provider = provider or GoogleSheets()
        self.lock = RLock()

    def connection(self, conn):
        return conn.execute('SELECT * FROM sheet_connection WHERE id=1').fetchone()

    def records(self, conn):
        return {kind: [json.loads(row['data']) for row in conn.execute('SELECT data FROM sheet_' + kind + ' ORDER BY id')]
                for kind in FIELDS}

    def status(self):
        with self.db_conn() as conn:
            connection = self.connection(conn)
            result = {'ok': True, **self.records(conn), 'connection': None,
                      'fields': FIELDS, 'required': REQUIRED,
                      'response_fields': form_responses.FIELDS, 'response_required': form_responses.REQUIRED,
                      'google_configured': bool(os.getenv('GOOGLE_APPLICATION_CREDENTIALS') or os.getenv('GOOGLE_SHEETS_USE_ADC') == 'true')}
            source = conn.execute('SELECT settings FROM sheet_source_preferences WHERE id=1').fetchone()
            result['source'] = json.loads(source['settings']) if source else None
            result['workspace_revision'] = conn.execute('SELECT revision FROM crm_workspace WHERE id=1').fetchone()[0]
            if connection:
                result['connection'] = {**json.loads(connection['settings']),
                    'revision': connection['revision'], 'last_success': connection['last_success'],
                    'last_attempt': connection['last_attempt'], 'last_error': connection['last_error'],
                    'report': json.loads(connection['last_result']) if connection['last_result'] else None}
            return result

    def provider_for(self, body):
        profile = body.get('public_source')
        if profile is None:
            return self.provider
        if body.get('mode', 'form_responses') != 'form_responses':
            raise SheetError('Existing-link access is supported only for one form-response worksheet.')
        return PublicFormSheet(parse_link(body.get('link')), profile, SheetError, MAX_BYTES)

    def settings(self, body):
        file_id = parse_link(body.get('link'))
        info = self.provider_for(body).inspect(file_id)
        mode = body.get('mode', 'tabs')
        if mode not in ('tabs', 'form_responses'):
            raise SheetError('Choose separate worksheets or a Google Form response worksheet.')
        if mode == 'form_responses':
            tab = body.get('response_tab')
            if not isinstance(tab, str) or tab not in {tab['name'] for tab in info['tabs']}:
                raise SheetError('Select the form response worksheet.')
            selected = {'leads': tab, 'queries': tab}
        else:
            selected = body.get('tabs')
        if not isinstance(selected, dict) or set(selected) != set(FIELDS) or any(not isinstance(selected.get(kind), str) for kind in FIELDS):
            raise SheetError('Select a Leads tab and a Queries tab.')
        if mode == 'tabs' and selected['leads'] == selected['queries']:
            raise SheetError('Choose separate Leads and Queries tabs.')
        available = {tab['name'] for tab in info['tabs']}
        if any(title not in available for title in selected.values()):
            raise SheetError('A selected tab no longer exists. Inspect the link again.')
        minutes = body.get('auto_minutes', 0)
        if isinstance(minutes, bool) or not isinstance(minutes, int) or minutes not in (0, 5, 15, 30, 60):
            raise SheetError('Automatic refresh must be off, 5, 15, 30 or 60 minutes.')
        mappings = body.get('mappings')
        if mode == 'tabs' and not isinstance(mappings, dict):
            raise SheetError('Map the source columns before previewing.')
        settings = {'link': info['url'], 'id': file_id, 'name': info['name'], 'tabs': selected,
                    'mappings': mappings or {}, 'auto_minutes': minutes}
        if body.get('public_source') is not None:
            settings['public_source'] = self.provider_for(body).profile
        if mode == 'form_responses':
            settings.update(mode=mode, response_tab=tab, response_mapping=body.get('response_mapping'),
                            form_link=form_responses.parse_form_link(body.get('form_link', ''), SheetError),
                            source_locale=info.get('locale', 'en_US'), source_timezone=info.get('timezone', 'Asia/Kolkata'))
        return settings

    def snapshot(self, settings, conn):
        if settings.get('mode') == 'form_responses':
            values = self.provider_for(settings).read(settings['id'], [settings['response_tab']])
            rows, errors, warnings = form_responses.convert(values[0], settings, validate_rows, SheetError)
            return {'settings': settings, 'rows': rows, 'errors': errors, 'warnings': warnings}
        values = self.provider.read(settings['id'], [settings['tabs'][kind] for kind in FIELDS])
        rows, errors = {}, []
        for kind, source in zip(FIELDS, values):
            rows[kind], rejected = validate_rows(kind, source, settings['mappings'].get(kind))
            errors.extend(rejected)
        # A rejected source lead cannot validate a query, even if an older copy exists.
        rejected_leads = {error['id'] for error in errors if error['kind'] == 'leads' and error['id']}
        lead_ids = {row['id'] for row in conn.execute('SELECT id FROM sheet_leads')} | {row['data']['id'] for row in rows['leads']}
        valid_queries = []
        for item in rows['queries']:
            lead_id = item['data']['lead_id']
            if lead_id not in lead_ids or lead_id in rejected_leads:
                errors.append({'kind': 'queries', 'row': item['row'], 'id': item['data']['id'],
                               'reason': 'Related Lead ID is missing or rejected. Import a valid lead with this exact ID first.'})
            else:
                valid_queries.append(item)
        rows['queries'] = valid_queries
        return {'settings': settings, 'rows': rows, 'errors': errors}

    def plan(self, conn, snapshot, apply=False):
        shared_leads = []
        form_mode = snapshot['settings'].get('mode') == 'form_responses'
        query_lead_ids = {item['data']['lead_id'] for item in snapshot['rows']['queries']}
        existing_shared_ids = {row['id'] for row in shared_workspace.read(conn)['values']['hk_v2_leads']} if form_mode and apply else set()
        report = {kind: {'added': 0, 'updated': 0, 'unchanged': 0,
                         'rejected': sum(error['kind'] == kind for error in snapshot['errors'])}
                  for kind in FIELDS}
        for kind in FIELDS:
            for item in snapshot['rows'][kind]:
                incoming = item['data']
                existing = conn.execute('SELECT data FROM sheet_' + kind + ' WHERE id=?', (incoming['id'],)).fetchone()
                old = json.loads(existing['data']) if existing else {}
                merged = {**old, **incoming}
                if kind == 'leads' and not merged.get('status'):
                    merged['status'] = 'Query Received' if form_mode and incoming['id'] in query_lead_ids else 'New Lead'
                if kind == 'queries' and not merged.get('status'):
                    merged['status'] = 'Open'
                if kind == 'queries' and form_mode:
                    if existing and merged != old:
                        merged['updated_date'] = self.now_iso()
                    elif not existing:
                        merged['updated_date'] = merged['created_date']
                if apply and kind == 'leads':
                    shared_row = dict(merged)
                    if form_mode and merged['id'] in existing_shared_ids:
                        shared_row.pop('status', None)  # Response forms do not own the CRM pipeline stage.
                    shared_leads.append(shared_row)
                changed = merged != old
                report[kind]['updated' if existing and changed else 'unchanged' if existing else 'added'] += 1
                if apply and changed:
                    if kind == 'leads':
                        conn.execute('INSERT INTO sheet_leads(id,data) VALUES (?,?) ON CONFLICT(id) DO UPDATE SET data=excluded.data', (merged['id'], json_text(merged)))
                    else:
                        conn.execute('INSERT INTO sheet_queries(id,lead_id,data) VALUES (?,?,?) ON CONFLICT(id) DO UPDATE SET lead_id=excluded.lead_id,data=excluded.data', (merged['id'], merged['lead_id'], json_text(merged)))
        if apply:
            shared_workspace.merge_sheet_leads(conn, shared_leads)
        return {'counts': report, 'errors': snapshot['errors'], 'warnings': snapshot.get('warnings', [])}

    def preview(self, body):
        with self.lock:
            settings = self.settings(body)
            with self.db_conn() as conn:
                snapshot = self.snapshot(settings, conn)
                report = self.plan(conn, snapshot)
                connection = self.connection(conn)
                revision = connection['revision'] if connection else 0
                token = secrets.token_urlsafe(32)
                now = datetime.now(timezone.utc).timestamp()
                conn.execute('DELETE FROM sheet_previews WHERE expires<?', (now,))
                conn.execute('INSERT INTO sheet_previews VALUES (?,?,?,?)', (token, now + 900, revision, json_text(snapshot)))
            return {'ok': True, 'token': token, 'report': report,
                    'preview': {kind: [row['data'] for row in snapshot['rows'][kind]][:10] for kind in FIELDS},
                    'name': settings['name']}

    def commit(self, token):
        with self.lock, self.db_conn() as conn:
            conn.execute('BEGIN IMMEDIATE')
            preview = conn.execute('SELECT * FROM sheet_previews WHERE token=?', (token,)).fetchone()
            if not preview or preview['expires'] < datetime.now(timezone.utc).timestamp():
                raise SheetError('Preview expired. Preview the data again before importing.', 409)
            connection = self.connection(conn)
            if preview['revision'] != (connection['revision'] if connection else 0):
                raise SheetError('Data changed after this preview. Preview again to review the latest changes.', 409)
            snapshot = json.loads(preview['snapshot'])
            report = self.plan(conn, snapshot, apply=True)
            now = self.now_iso()
            conn.execute('''INSERT INTO sheet_connection(id,settings,revision,last_success,last_attempt,last_result,last_attempt_epoch)
                VALUES(1,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET settings=excluded.settings,
                revision=excluded.revision,last_success=excluded.last_success,last_attempt=excluded.last_attempt,
                last_result=excluded.last_result,last_error=NULL,last_attempt_epoch=excluded.last_attempt_epoch''',
                (json_text(snapshot['settings']), preview['revision'] + 1, now, now, json_text(report), datetime.now(timezone.utc).timestamp()))
            conn.execute('DELETE FROM sheet_previews')
            settings = snapshot['settings']
            if settings.get('mode') == 'form_responses':
                source = {key: settings[key] for key in ('link', 'form_link', 'mode')}
                conn.execute('INSERT INTO sheet_source_preferences VALUES (1,?) ON CONFLICT(id) DO UPDATE SET settings=excluded.settings', (json_text(source),))
        return self.status()

    def sync(self, due=False):
        with self.lock:
            with self.db_conn() as conn:
                connection = self.connection(conn)
            if not connection:
                if due:
                    return None
                raise SheetError('Connect, preview and import a spreadsheet first.', 409)
            settings = json.loads(connection['settings'])
            now_epoch = datetime.now(timezone.utc).timestamp()
            if due and (not settings['auto_minutes'] or now_epoch - connection['last_attempt_epoch'] < settings['auto_minutes'] * 60):
                return None
            try:
                # Revalidate access, format and selected tabs on every run.
                current = self.settings(settings)
                with self.db_conn() as conn:
                    snapshot = self.snapshot(current, conn)
                    conn.execute('BEGIN IMMEDIATE')
                    report = self.plan(conn, snapshot, apply=True)
                    now = self.now_iso()
                    conn.execute('UPDATE sheet_connection SET settings=?,revision=revision+1,last_success=?,last_attempt=?,last_error=NULL,last_result=?,last_attempt_epoch=? WHERE id=1',
                                 (json_text(current), now, now, json_text(report), now_epoch))
            except Exception as exc:
                error = exc if isinstance(exc, SheetError) else SheetError('Synchronization failed. The last successfully saved data is retained.', 502)
                with self.db_conn() as conn:
                    conn.execute('UPDATE sheet_connection SET last_attempt=?,last_attempt_epoch=?,last_error=? WHERE id=1', (self.now_iso(), now_epoch, str(error)))
                raise error from None
        return self.status()

    def auto_sync(self):
        try:
            self.sync(due=True)
        except SheetError:
            pass  # Persisted error is visible in the dashboard; retry next interval.


def export_workbook(payload):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    if not isinstance(payload, dict):
        raise SheetError('Provide Leads and Queries records for export.')
    workbook = Workbook()
    workbook.remove(workbook.active)
    for kind, fields in FIELDS.items():
        rows = payload.get(kind)
        if not isinstance(rows, list) or len(rows) > MAX_ROWS or any(not isinstance(row, dict) for row in rows):
            raise SheetError('Export up to 20,000 records per worksheet.')
        sheet = workbook.create_sheet(kind.title())
        sheet.append([label for _, label in fields])
        for record in rows:
            values = []
            for field, _ in fields:
                value = record.get(field, '')
                if not isinstance(value, (str, int, float, bool)) and value is not None:
                    raise SheetError('Export fields must contain plain values.')
                if field in DATE_FIELDS and value:
                    try:
                        parsed = datetime.fromisoformat(normalize_date(str(value)))
                        value = parsed.replace(tzinfo=None) if parsed.tzinfo is None else parsed.astimezone(timezone.utc).replace(tzinfo=None)
                    except ValueError:
                        # Keep a legacy display date visible rather than guessing its locale.
                        value = str(value)
                elif field not in ('quote', 'order'):
                    value = '' if value is None else str(value)
                values.append(value)
            sheet.append(values)
            for cell, (field, _) in zip(sheet[sheet.max_row], fields):
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Prevent spreadsheet formula injection, including IDs beginning '='.
                    cell.number_format = '@'
                elif field in DATE_FIELDS:
                    cell.number_format = 'yyyy-mm-dd hh:mm' if 'T' in str(record.get(field, '')) else 'yyyy-mm-dd'
                elif field in ('quote', 'order'):
                    cell.number_format = '#,##0.00'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='172A52')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 28
        for index, (field, label) in enumerate(fields, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 44 if field in ('details', 'requirement', 'remark') else max(18, len(label) + 3)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def register(app, db_conn, now_iso, authorize_admin):
    service = SheetSync(db_conn, now_iso)
    routes = Blueprint('spreadsheet', __name__, url_prefix='/api/spreadsheet')

    @routes.before_request
    def same_origin():
        if request.method == 'POST':
            if not crm_request_origin_allowed(request):
                raise SheetError('Use the dashboard on this server to manage spreadsheet data.', 403)
            if request.headers.get('X-CRM-Request') != '1' or not request.is_json:
                raise SheetError('Send a JSON request from the CRM dashboard.', 403)
            if request.content_length and request.content_length > MAX_BYTES:
                raise SheetError('Request exceeds the 12 MB limit.', 413)
            if request.endpoint.rsplit('.', 1)[-1] in {'inspect', 'save_source', 'columns', 'preview', 'commit'}:
                authorize_admin(body())

    @routes.after_request
    def private_response(response):
        response.headers['Cache-Control'] = 'no-store, private'
        return response

    @routes.errorhandler(SheetError)
    def handle_error(error):
        return jsonify(ok=False, error=str(error)), error.status

    @routes.errorhandler(sqlite3.Error)
    def database_error(error):
        return jsonify(ok=False, error='The database could not save this operation. No partial import was saved; the last successful data is retained. Check disk space and database access, then retry.'), 503

    def body():
        value = request.get_json(silent=True)
        if not isinstance(value, dict):
            raise SheetError('Provide a JSON object.')
        return value

    @routes.get('/status')
    def status():
        return jsonify(service.status())

    @routes.post('/inspect')
    def inspect():
        data = body()
        return jsonify(ok=True, **service.provider_for(data).inspect(parse_link(data.get('link'))))

    @routes.post('/source')
    def save_source():
        data = body()
        file_id = parse_link(data.get('link'))
        form_link = form_responses.parse_form_link(data.get('form_link', ''), SheetError)
        mode = data.get('mode', 'form_responses')
        if mode not in ('tabs', 'form_responses'):
            raise SheetError('Invalid spreadsheet source mode.')
        source = {'link': 'https://docs.google.com/spreadsheets/d/' + file_id + '/edit',
                  'form_link': form_link, 'mode': mode}
        with db_conn() as conn:
            conn.execute('INSERT INTO sheet_source_preferences VALUES (1,?) ON CONFLICT(id) DO UPDATE SET settings=excluded.settings', (json_text(source),))
        return jsonify(service.status())

    @routes.post('/columns')
    def columns():
        data = body()
        file_id = parse_link(data.get('link'))
        provider = service.provider_for(data)
        info = provider.inspect(file_id)
        titles = data.get('tabs')
        if not isinstance(titles, list) or len(titles) not in (1, 2) or any(title not in [tab['name'] for tab in info['tabs']] for title in titles):
            raise SheetError('Select valid Leads and Queries tabs.')
        values = provider.read(file_id, titles, headers_only=True)
        return jsonify(ok=True, sheets=[{'headers': headers_for(rows), 'sample': rows[1:6]} for rows in values])

    @routes.post('/preview')
    def preview():
        return jsonify(service.preview(body()))

    @routes.post('/import')
    def commit():
        token = body().get('token')
        if not isinstance(token, str) or len(token) > 100:
            raise SheetError('Preview the spreadsheet first.')
        return jsonify(service.commit(token))

    @routes.post('/sync')
    def sync():
        return jsonify(service.sync())

    @routes.post('/export')
    def export():
        output = export_workbook(body())
        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True, download_name='HOSPkart_Leads_Queries_' + now_iso()[:10] + '.xlsx')

    app.register_blueprint(routes)
    return service
